import base64
import hashlib
import json
import markdown
import re
import shlex
import signal
import socket
import ssl
import subprocess
import time
import threading
import http.cookiejar
import urllib.error
import urllib.parse
import urllib.request
import uuid
import webbrowser
import zipfile
import unicodedata
import requests
from html import escape as html_escape
from io import BytesIO
from xml.sax.saxutils import escape as xml_escape
from cryptography.hazmat.primitives import padding as sym_padding, serialization
from cryptography.hazmat.primitives.asymmetric import padding as asym_padding
from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes
from flask import Flask, render_template, request, jsonify, send_file, abort
import paramiko
import pymysql
import os
from datetime import datetime, timedelta
import glob
from pathlib import Path
try:
    import certifi
except ImportError:
    certifi = None

from config import LOG_MAPPING, SSH_CONFIG, DB_CONFIG, DOCS_BASE_PATH, SKILLS_BASE_PATH, PORT, _manager
import photo_classify

app = Flask(__name__)
_ORACLE_CLIENT_INITIALIZED = False
LOCAL_LOG_DIR = "/Users/wangpengjing/sbyProject/birds_tools/logs"
SERVICE_VUE_PROCESS = None
SERVICE_VUE_LOGS = []
SERVICE_VUE_LOCK = threading.Lock()
SERVICE_VUE_MAX_LOG_LINES = 1200

CACHE_TTL_SECONDS = 300
_CACHE_LOCK = threading.RLock()
_CACHE = {}
_PROTECTED_WILDLIFE_CACHE_KEY = "protected_wildlife_catalog"
_BIRD_QUERY_CACHE_PREFIX = "bird-query:"
_LATEST_MEMORY_BATCH = None


def _cache_get(key):
    now = time.monotonic()
    with _CACHE_LOCK:
        item = _CACHE.get(key)
        if not item:
            return None
        expires_at, value = item
        if expires_at <= now:
            _CACHE.pop(key, None)
            return None
        return value


def _cache_set(key, value, ttl=CACHE_TTL_SECONDS):
    with _CACHE_LOCK:
        _CACHE[key] = (time.monotonic() + ttl, value)
    return value


def _load_protected_wildlife_from_workbook():
    """从项目内置 Excel 读取保护动物名录，供无数据库模式和数据库故障回退使用。"""
    try:
        import openpyxl
        workbook_path = Path(__file__).resolve().parent / "国家重点保护野生动物名录.xlsx"
        workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)
        sheet = workbook["国家重点保护动物"]
        rows = sheet.iter_rows(values_only=True)
        for _ in range(7):
            next(rows, None)
        result = []
        for row in rows:
            if not row or not row[0]:
                continue
            result.append({
                "chinese_name": str(row[1] or "").strip(),
                "latin_name": str(row[2] or "").strip(),
                "taxonomy": str(row[3] or "").strip(),
                "protection_level": str(row[4] or "").strip(),
            })
        workbook.close()
        return result
    except Exception as exc:
        print(f"[Cache] 读取保护动物名录 Excel 失败: {exc}")
        return []


def load_protected_wildlife_cache(force=False):
    """启动时及缓存过期后加载保护动物名录，缓存五分钟。"""
    if not force:
        cached = _cache_get(_PROTECTED_WILDLIFE_CACHE_KEY)
        if cached is not None:
            return cached
    rows = []
    if _manager.db_enabled:
        conn = None
        try:
            conn = get_db()
            with conn.cursor() as cur:
                cur.execute("""
                    SELECT chinese_name, latin_name, taxonomy, protection_level
                    FROM protected_wildlife_catalog ORDER BY id ASC
                """)
                rows = dict_rows(cur)
        except Exception as exc:
            print(f"[Cache] 从数据库加载保护动物名录失败，改用 Excel: {exc}")
        finally:
            if conn:
                conn.close()
    if not rows:
        rows = _load_protected_wildlife_from_workbook()
    return _cache_set(_PROTECTED_WILDLIFE_CACHE_KEY, rows)


def get_protected_wildlife_level(names):
    names = set(names or [])
    return {
        row["chinese_name"]: row.get("protection_level")
        for row in load_protected_wildlife_cache()
        if row.get("chinese_name") in names
    }


def _query_cache_key(prefix, payload):
    return prefix + hashlib.sha256(json.dumps(payload, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def warmup_caches():
    """应用启动时预热保护动物名录缓存。"""
    load_protected_wildlife_cache(force=True)

# 主页工具卡片定义
TOOLS = [
    {
        "id": "config-manager",
        "name": "配置管理",
        "description": "管理应用配置参数，支持在线修改 SSH、数据库、路径等配置。",
        "icon": "fa-cogs",
        "color": "#8b5cf6"
    },
    {
        "id": "bird-records",
        "name": "鸟种记录",
        "description": "维护常用观鸟地点，点击地点后查询最近一天的鸟种记录。",
        "icon": "fa-dove",
        "color": "#16a34a"
    },
    {
        "id": "ai-bird-chat",
        "name": "AI 观鸟问答",
        "description": "对接 DeepSeek，用自然语言查询鸟种记录并自动总结。",
        "icon": "fa-comments",
        "color": "#2563eb"
    },
    {
        "id": "photo-classify",
        "name": "照片分类管理",
        "description": "扫描照片、按拍摄日期分组，并复制或移动到日期目录。",
        "icon": "fa-images",
        "color": "#e91e63"
    },
    {
        "id": "bird-navigation",
        "name": "小鸟导航",
        "description": "按保护鸟种和省市区查询观鸟记录，并跳转地图查看观测位置。",
        "icon": "fa-binoculars",
        "color": "#0f766e"
    }
]

@app.route('/')
def index():
    return render_template('index.html', tools=get_ordered_tools())

def get_ordered_tools():
    tool_map = {tool["id"]: tool for tool in TOOLS}
    configured_order = _manager.get_all().get("home_tool_order", [])
    enabled = _manager.get_all().get("home_tool_enabled") or {}
    ordered = []
    seen = set()
    for tool_id in configured_order:
        tool = tool_map.get(tool_id)
        if tool and tool_id not in seen and enabled.get(tool_id, tool_id != "photo-classify") is not False:
            ordered.append(tool)
            seen.add(tool_id)
    for tool in TOOLS:
        if tool["id"] not in seen and enabled.get(tool["id"], tool["id"] != "photo-classify") is not False:
            ordered.append(tool)
    return ordered

TOOL_VIEWER_CONFIG = {
    "doc-query": {
        "template": "tool_md_viewer.html",
        "api_tree": "/api/docs/tree",
        "api_view": "/api/docs/view",
        "sidebar_root": "feature-doc",
        "tool_title": "服务商文档查询",
        "tool_icon": "fa-book",
        "tool_color": "#0891b2",
        "tool_description": "浏览和查看 feature-doc 目录下的设计文档",
    },
    "skill-query": {
        "template": "tool_md_viewer.html",
        "api_tree": "/api/skills/tree",
        "api_view": "/api/skills/view",
        "sidebar_root": "skills",
        "tool_title": "服务商代码查询",
        "tool_icon": "fa-code",
        "tool_color": "#7c3aed",
        "tool_description": "浏览 .agents/skills 目录下的 SKILL 文档",
    },
}

# ---- MySQL / Bookmarks ----

def get_db():
    """获取数据库连接"""
    if not _manager.db_enabled:
        raise RuntimeError("数据库未启用，当前功能使用五分钟缓存；需要持久化数据时请在配置管理中启用数据库")
    return pymysql.connect(**_manager.db_config)

def init_bookmarks_table():
    """初始化收藏表"""
    if not _manager.db_enabled:
        return
    conn = None
    try:
        conn = get_db()
        with conn.cursor() as cur:
            cur.execute("""
                CREATE TABLE IF NOT EXISTS bookmarks (
                    id INT AUTO_INCREMENT PRIMARY KEY,
                    doc_name VARCHAR(500) NOT NULL,
                    doc_path VARCHAR(1000) NOT NULL UNIQUE,
                    source VARCHAR(50) DEFAULT 'doc-query',
                    created_at DATETIME DEFAULT CURRENT_TIMESTAMP
                ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
            """)
            try:
                cur.execute("ALTER TABLE bookmarks ADD COLUMN source VARCHAR(50) DEFAULT 'doc-query'")
            except Exception:
                pass  # 兼容旧表
        conn.commit()
        conn.close()
    except Exception as e:
        print(f"[DB] 初始化收藏表失败: {e}")

@app.route('/api/bookmarks', methods=['GET'])
def list_bookmarks():
    """获取所有收藏的文档路径"""
    try:
        conn = get_db()
        with conn.cursor() as cur:
            cur.execute("SELECT doc_path FROM bookmarks")
            rows = cur.fetchall()
        conn.close()
        return jsonify([row[0] for row in rows])
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/api/bookmarks/toggle', methods=['POST'])
def toggle_bookmark():
    """切换收藏状态：传 doc_path，返回新的状态"""
    data = request.json
    doc_path = data.get("doc_path", "").strip()
    doc_name = data.get("doc_name", "").strip()
    source = data.get("source", "doc-query")
    if not doc_path:
        return jsonify({"error": "缺少 doc_path"}), 400
    try:
        conn = get_db()
        with conn.cursor() as cur:
            cur.execute("SELECT id FROM bookmarks WHERE doc_path = %s", (doc_path,))
            row = cur.fetchone()
            if row:
                cur.execute("DELETE FROM bookmarks WHERE doc_path = %s", (doc_path,))
                conn.commit()
                conn.close()
                return jsonify({"bookmarked": False, "doc_path": doc_path})
            else:
                cur.execute(
                    "INSERT INTO bookmarks (doc_name, doc_path, source) VALUES (%s, %s, %s)",
                    (doc_name or doc_path, doc_path, source)
                )
                conn.commit()
                conn.close()
                return jsonify({"bookmarked": True, "doc_path": doc_path})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/api/bookmarks/detail', methods=['GET'])
def bookmarks_detail():
    """返回收藏详情（含目录信息）"""
    try:
        conn = get_db()
        with conn.cursor() as cur:
            cur.execute("SELECT doc_name, doc_path, source, created_at FROM bookmarks ORDER BY created_at DESC")
            rows = cur.fetchall()
        conn.close()
        result = []
        for row in rows:
            doc_name, doc_path, source, created_at = row
            parts = doc_path.replace('\\', '/').split('/')
            directory = '/'.join(parts[:-1]) if len(parts) > 1 else ''
            result.append({
                "doc_name": doc_name,
                "doc_path": doc_path,
                "source": source or "doc-query",
                "directory": directory,
                "created_at": created_at.isoformat() if created_at else ""
            })
        return jsonify(result)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/tools/bookmarks')
def bookmarks_page():
    return render_template('tool_bookmarks.html', tools=TOOLS)

# ---- 配置中心 ----

@app.route('/api/config', methods=['GET'])
def get_config():
    """获取全部配置"""
    return jsonify(_manager.get_all())

@app.route('/api/config', methods=['PUT'])
def update_config():
    """更新配置"""
    data = request.json
    if not data:
        return jsonify({"error": "请求体不能为空"}), 400
    try:
        _manager.update(data)
        # 数据库开关或连接信息变化后，让名录缓存按新配置重新加载。
        with _CACHE_LOCK:
            _CACHE.pop(_PROTECTED_WILDLIFE_CACHE_KEY, None)
        warmup_caches()
        return jsonify({"status": "ok", "message": "配置已更新，部分配置需重启应用生效"})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/tools/config-manager')
def config_manager_page():
    return render_template('tool_config.html', tools=TOOLS, PORT=PORT, home_tools=TOOLS)

# ---- Git 分支操作 ----

DEFAULT_GIT_PROJECTS = [
    {
        "id": "tax_pay_ext",
        "name": "tax_pay_ext",
        "project_path": "/Users/wangpengjing/sbyProject/tax_pay_ext",
        "branches": ["test", "release"],
    },
    {
        "id": "risk-cloud",
        "name": "risk-cloud",
        "project_path": "/Users/wangpengjing/sbyProject/risk-cloud",
        "branches": ["main", "test"],
    },
]
DEFAULT_GIT_BRANCHES = DEFAULT_GIT_PROJECTS[0]["branches"]
DEFAULT_GIT_PROJECT_PATH = DEFAULT_GIT_PROJECTS[0]["project_path"]

def normalize_git_branches(branches, default_branches=None):
    result = []
    for branch in branches or []:
        branch = str(branch or "").strip()
        if branch and branch not in result:
            result.append(branch)
    for branch in default_branches or DEFAULT_GIT_BRANCHES:
        if branch not in result:
            result.insert(len(result), branch)
    return result

def normalize_bool(value, default=False):
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return bool(value)
    text = str(value).strip().lower()
    if text in {"1", "true", "yes", "on"}:
        return True
    if text in {"0", "false", "no", "off", ""}:
        return False
    return default

def default_git_project_map():
    return {project["id"]: dict(project) for project in DEFAULT_GIT_PROJECTS}

def get_git_branch_manager_config():
    cfg = _manager.get_all().get("git_branch_manager", {})
    default_projects = default_git_project_map()
    projects = []
    seen = set()

    configured_projects = cfg.get("projects")
    if configured_projects:
        for project in configured_projects:
            project_id = str(project.get("id") or "").strip()
            if not project_id or project_id in seen:
                continue
            defaults = default_projects.get(project_id, {})
            default_branches = defaults.get("branches") or DEFAULT_GIT_BRANCHES
            projects.append({
                "id": project_id,
                "name": (project.get("name") or defaults.get("name") or project_id).strip(),
                "project_path": (project.get("project_path") or defaults.get("project_path") or "").strip(),
                "branches": normalize_git_branches(project.get("branches") or default_branches, default_branches),
            })
            seen.add(project_id)
    else:
        projects.append({
            "id": "tax_pay_ext",
            "name": "tax_pay_ext",
            "project_path": (cfg.get("project_path") or DEFAULT_GIT_PROJECT_PATH).strip(),
            "branches": normalize_git_branches(cfg.get("branches") or DEFAULT_GIT_BRANCHES, DEFAULT_GIT_BRANCHES),
        })
        seen.add("tax_pay_ext")

    for project_id, defaults in default_projects.items():
        if project_id not in seen:
            projects.append({
                "id": defaults["id"],
                "name": defaults["name"],
                "project_path": defaults["project_path"],
                "branches": normalize_git_branches(defaults["branches"], defaults["branches"]),
            })

    current_project = (cfg.get("current_project") or "").strip()
    if not current_project or current_project not in {project["id"] for project in projects}:
        current_project = projects[0]["id"]

    return {
        "current_project": current_project,
        "projects": projects,
        "auto_merge_release_after_switch": normalize_bool(cfg.get("auto_merge_release_after_switch"), False),
    }

def get_git_project_config(project_id=None):
    cfg = get_git_branch_manager_config()
    project_id = (project_id or cfg["current_project"] or "").strip()
    for project in cfg["projects"]:
        if project["id"] == project_id:
            return project, cfg["projects"]
    raise ValueError(f"未知项目: {project_id}")

def save_git_project_config(project, projects, extra_config=None):
    next_projects = []
    updated = False
    for item in projects:
        if item["id"] == project["id"]:
            next_projects.append(project)
            updated = True
        else:
            next_projects.append(item)
    if not updated:
        next_projects.append(project)
    git_branch_manager = {"current_project": project["id"], "projects": next_projects}
    if extra_config:
        git_branch_manager.update(extra_config)
    _manager.update({"git_branch_manager": git_branch_manager})

def validate_git_branch_name(branch):
    branch = (branch or "").strip()
    if not branch:
        raise ValueError("分支名不能为空")
    if len(branch) > 120:
        raise ValueError("分支名过长")
    invalid = (
        branch.startswith("-") or branch.startswith("/") or branch.endswith("/") or
        branch.endswith(".") or branch.endswith(".lock") or
        ".." in branch or "@{" in branch or "\\" in branch or
        any(ch.isspace() or ord(ch) < 32 for ch in branch)
    )
    if invalid or not re.fullmatch(r"[A-Za-z0-9._/-]+", branch):
        raise ValueError("分支名只能包含字母、数字、点、下划线、横线和斜杠，且不能使用 Git 非法格式")
    return branch

def run_git(project_path, args, timeout=120):
    completed = subprocess.run(
        ["git", "-C", project_path] + args,
        capture_output=True,
        text=True,
        timeout=timeout,
    )
    output = (completed.stdout or "").strip()
    error = (completed.stderr or "").strip()
    return completed.returncode, output, error

def require_git_project(project_path):
    if not os.path.isdir(project_path):
        raise FileNotFoundError(f"项目目录不存在: {project_path}")
    code, out, err = run_git(project_path, ["rev-parse", "--is-inside-work-tree"], timeout=20)
    if code != 0 or out.strip() != "true":
        raise ValueError(f"目标目录不是 Git 仓库: {project_path}")

def git_ref_exists(project_path, ref):
    code, _, _ = run_git(project_path, ["show-ref", "--verify", "--quiet", ref], timeout=20)
    return code == 0

def get_git_status_payload(project_id=None):
    project, projects = get_git_project_config(project_id)
    cfg = get_git_branch_manager_config()
    project_path = project["project_path"]
    require_git_project(project_path)
    code, branch, err = run_git(project_path, ["rev-parse", "--abbrev-ref", "HEAD"], timeout=20)
    if code != 0:
        raise RuntimeError(err or "获取当前分支失败")
    code, status, err = run_git(project_path, ["status", "--porcelain"], timeout=30)
    if code != 0:
        raise RuntimeError(err or "获取工作区状态失败")
    code, remote, _ = run_git(project_path, ["remote", "get-url", "origin"], timeout=20)
    return {
        "project_id": project["id"],
        "project_name": project["name"],
        "project_path": project_path,
        "projects": projects,
        "branches": project["branches"],
        "current_branch": branch,
        "dirty": bool(status.strip()),
        "dirty_count": len([line for line in status.splitlines() if line.strip()]),
        "remote": remote if code == 0 else "",
        "auto_merge_release_after_switch": cfg["auto_merge_release_after_switch"],
    }

@app.route('/api/git-branch-manager/status', methods=['GET'])
def api_git_branch_manager_status():
    try:
        return jsonify(get_git_status_payload(request.args.get("project_id")))
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/api/git-branch-manager/branches', methods=['POST'])
def api_git_branch_manager_add_branch():
    data = request.get_json(silent=True) or {}
    try:
        branch = validate_git_branch_name(data.get("branch"))
        project, projects = get_git_project_config(data.get("project_id"))
        project["branches"] = normalize_git_branches(project["branches"] + [branch], project["branches"])
        save_git_project_config(project, projects)
        return jsonify({"status": "ok", "branches": project["branches"], "project_id": project["id"]})
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/api/git-branch-manager/preferences', methods=['POST'])
def api_git_branch_manager_preferences():
    data = request.get_json(silent=True) or {}
    try:
        auto_merge_release_after_switch = normalize_bool(data.get("auto_merge_release_after_switch"), False)
        _manager.update({
            "git_branch_manager": {
                "auto_merge_release_after_switch": auto_merge_release_after_switch,
            }
        })
        return jsonify({
            "status": "ok",
            "auto_merge_release_after_switch": auto_merge_release_after_switch,
        })
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/api/git-branch-manager/switch', methods=['POST'])
def api_git_branch_manager_switch():
    data = request.get_json(silent=True) or {}
    try:
        branch = validate_git_branch_name(data.get("branch"))
        project, projects = get_git_project_config(data.get("project_id"))
        auto_merge_release_after_switch = normalize_bool(
            data.get("auto_merge_release_after_switch"),
            get_git_branch_manager_config()["auto_merge_release_after_switch"],
        )
        project_path = project["project_path"]
        require_git_project(project_path)

        steps = []
        for args, label, timeout in (
            (["fetch", "origin", "--prune"], "拉取远端引用", 180),
            (["switch", branch], "切换本地分支", 120),
        ):
            if label == "切换本地分支" and not git_ref_exists(project_path, f"refs/heads/{branch}"):
                if git_ref_exists(project_path, f"refs/remotes/origin/{branch}"):
                    args = ["switch", "--track", "-c", branch, f"origin/{branch}"]
                else:
                    return jsonify({"error": f"未找到本地分支或远端 origin/{branch}"}), 404
            code, out, err = run_git(project_path, args, timeout=timeout)
            steps.append({"step": label, "command": "git " + " ".join(args), "output": out, "error": err})
            if code != 0:
                return jsonify({"error": err or out or f"{label}失败", "steps": steps}), 500

        code, upstream, _ = run_git(project_path, ["rev-parse", "--abbrev-ref", "--symbolic-full-name", "@{u}"], timeout=20)
        if code == 0 and upstream:
            pull_args = ["pull", "--ff-only"]
        else:
            pull_args = ["pull", "--ff-only", "origin", branch]
        code, out, err = run_git(project_path, pull_args, timeout=180)
        steps.append({"step": "拉取最新代码", "command": "git " + " ".join(pull_args), "output": out, "error": err})
        if code != 0:
            return jsonify({"error": err or out or "拉取最新代码失败", "steps": steps}), 500

        if auto_merge_release_after_switch and branch != "release":
            if git_ref_exists(project_path, "refs/remotes/origin/release"):
                merge_source = "origin/release"
            elif git_ref_exists(project_path, "refs/heads/release"):
                merge_source = "release"
            else:
                return jsonify({"error": "未找到本地 release 分支或远端 origin/release", "steps": steps}), 404
            merge_args = ["merge", "--no-edit", merge_source]
            code, out, err = run_git(project_path, merge_args, timeout=180)
            steps.append({
                "step": "合并 release 到当前分支",
                "command": "git " + " ".join(merge_args),
                "output": out,
                "error": err,
            })
            if code != 0:
                return jsonify({"error": err or out or "合并 release 失败", "steps": steps}), 500

        project["branches"] = normalize_git_branches(project["branches"] + [branch], project["branches"])
        save_git_project_config(
            project,
            projects,
            {"auto_merge_release_after_switch": auto_merge_release_after_switch},
        )
        status_payload = get_git_status_payload(project["id"])
        message = f"已切换到 {branch} 并拉取最新代码"
        if auto_merge_release_after_switch and branch != "release":
            message += "，已自动合并 release"
        return jsonify({"status": "ok", "message": message, "steps": steps, **status_payload})
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except subprocess.TimeoutExpired:
        return jsonify({"error": "Git 命令执行超时"}), 504
    except Exception as e:
        return jsonify({"error": str(e)}), 500

# ---- AI 图片生成 ----

def get_image_generation_config():
    cfg = _manager.get_all().get("image_generation", {})
    return {
        "model": (cfg.get("model") or "gpt-image-2").strip(),
        "api_key": (cfg.get("api_key") or "").strip(),
        "base_url": (cfg.get("base_url") or "https://api.0-0.pro/v1").strip(),
    }

def build_image_generation_url(base_url):
    base_url = (base_url or "").strip().rstrip("/")
    if not base_url:
        raise ValueError("请先在配置中心填写图片生成调用地址")
    if base_url.endswith("/images/generations"):
        return base_url
    return f"{base_url}/images/generations"

def extract_image_result(response_json):
    images = response_json.get("data")
    if not isinstance(images, list) or not images:
        raise ValueError("图片生成接口未返回 data 图片数据")
    image = images[0] or {}
    b64_json = image.get("b64_json") or image.get("base64") or image.get("image")
    image_url = image.get("url")
    data_url = ""
    if b64_json:
        data_url = b64_json if str(b64_json).startswith("data:") else f"data:image/png;base64,{b64_json}"
    return {
        "image_url": image_url or "",
        "data_url": data_url,
        "revised_prompt": image.get("revised_prompt") or response_json.get("revised_prompt") or "",
    }

@app.route('/api/image-generation/generate', methods=['POST'])
def api_image_generation_generate():
    data = request.get_json(silent=True) or {}
    prompt = (data.get("prompt") or "").strip()
    if not prompt:
        return jsonify({"error": "请输入图片需求"}), 400

    cfg = get_image_generation_config()
    if not cfg["api_key"]:
        return jsonify({"error": "请先在配置中心填写图片生成 API Key"}), 400
    if not cfg["model"]:
        return jsonify({"error": "请先在配置中心填写图片生成模型"}), 400

    size = (data.get("size") or "1024x1024").strip()
    quality = (data.get("quality") or "auto").strip()
    payload = {
        "model": cfg["model"],
        "prompt": prompt,
        "n": 1,
        "size": size,
    }
    if quality and quality != "auto":
        payload["quality"] = quality

    try:
        resp = requests.post(
            build_image_generation_url(cfg["base_url"]),
            headers={
                "Authorization": f"Bearer {cfg['api_key']}",
                "Content-Type": "application/json",
            },
            json=payload,
            timeout=120,
        )
        try:
            response_json = resp.json()
        except ValueError:
            response_json = {"error": {"message": resp.text[:500]}}

        if not resp.ok:
            upstream_error = response_json.get("error") if isinstance(response_json, dict) else None
            message = upstream_error.get("message") if isinstance(upstream_error, dict) else None
            return jsonify({"error": message or f"图片生成接口返回 {resp.status_code}"}), 502

        result = extract_image_result(response_json)
        if not result["data_url"] and not result["image_url"]:
            return jsonify({"error": "图片生成成功但未返回可展示的图片内容"}), 502
        return jsonify({
            "status": "ok",
            "model": cfg["model"],
            "size": size,
            "quality": quality,
            **result,
        })
    except requests.Timeout:
        return jsonify({"error": "图片生成请求超时，请稍后重试"}), 504
    except requests.RequestException as e:
        return jsonify({"error": f"图片生成请求失败: {str(e)}"}), 502
    except ValueError as e:
        return jsonify({"error": str(e)}), 502
    except Exception as e:
        return jsonify({"error": str(e)}), 500

# ---- 网络配置 ----

NETWORK_SERVICE = "USB 10/100/1000 LAN"

def run_command(command):
    try:
        subprocess.run(command, check=True, shell=True, capture_output=True, text=True, timeout=15)
        return True, ""
    except subprocess.CalledProcessError as e:
        return False, e.stderr.strip()
    except subprocess.TimeoutExpired:
        return False, "命令执行超时"
    except FileNotFoundError:
        return False, "未找到 networksetup 命令"

def set_to_dhcp(service_name):
    ok, err = run_command(f'networksetup -setdhcp "{service_name}"')
    if not ok:
        raise RuntimeError(err or "设置 DHCP 失败")
def set_to_manual(service_name, ip, subnet, router):
    ok, err = run_command(f'networksetup -setmanual "{service_name}" {ip} {subnet} {router}')
    if not ok:
        raise RuntimeError(err or "设置静态 IP 失败")

def modify_config(mode, ip="", subnet="", router=""):
    if mode == "dhcp":
        set_to_dhcp("Wi-Fi")
    elif mode == "manual":
        set_to_manual("Wi-Fi", ip, subnet, router)

def get_network_info():
    ok, out = run_command(f'networksetup -getinfo "{NETWORK_SERVICE}"')
    if not ok:
        raise RuntimeError(out or "获取网络配置失败")
    result = {"dhcp": False, "ip": "", "subnet": "", "router": ""}
    for line in out.split("\n"):
        line = line.strip()
        if line.startswith("DHCP Configuration"):
            result["dhcp"] = True
        elif line.startswith("Manual Configuration"):
            result["dhcp"] = False
        elif line.startswith("IP address:"):
            result["ip"] = line.split(":", 1)[1].strip()
        elif line.startswith("Subnet mask:"):
            result["subnet"] = line.split(":", 1)[1].strip()
        elif line.startswith("Router:"):
            val = line.split(":", 1)[1].strip()
            if val and val != "none":
                result["router"] = val
    return result

@app.route('/api/network', methods=['GET'])
def api_get_network():
    try:
        info = get_network_info()
        info["service"] = NETWORK_SERVICE
        return jsonify(info)
    except RuntimeError as e:
        return jsonify({"error": str(e)}), 500

@app.route('/api/network', methods=['POST'])
def api_set_network():

    data = request.json
    mode = data.get("mode", "").strip()
    print(f"测试..."+mode)

    if mode not in ("dhcp", "manual"):
        return jsonify({"error": "mode 必须为 dhcp 或 manual"}), 400
    ip = data.get("ip", "").strip()
    subnet = data.get("subnet", "").strip()
    router = data.get("router", "").strip()
    if mode == "manual" and (not ip or not subnet or not router):
        return jsonify({"error": "手动模式需要填写 IP、子网掩码、路由器"}), 400
    try:
        modify_config(mode, ip, subnet, router)
        msg = "已切换为 DHCP 模式" if mode == "dhcp" else f"已设置为静态 IP: {ip}"
        return jsonify({"status": "ok", "message": msg})
    except RuntimeError as e:
        return jsonify({"error": str(e)}), 500

# ---- 鸟种记录 ----

BIRDREPORT_API = "https://api.birdreport.cn/"
BIRDREPORT_PUBLIC_KEY = (
    "MIGfMA0GCSqGSIb3DQEBAQUAA4GNADCBiQKBgQCvxXa98E1uWXnBzXkS2yHUfnBM6n3PCwLdfIox03T91joBvjtoDqiQ5x3t"
    "TOfpHs3LtiqMMEafls6b0YWtgB1dse1W5m+FpeusVkCOkQxB4SZDH6tuerIknnmB/Hsq5wgEkIvO5Pff9biig6AyoAkdWp"
    "Sek/1/B7zYIepYY0lxKQIDAQAB"
)
BIRDREPORT_AES_KEY = b"C8EB5514AF5ADDB94B2207B08C66601C"
BIRDREPORT_AES_IV = b"55DD79C6F04E1A67"
BIRDREPORT_COOKIE_JAR = http.cookiejar.CookieJar()
BIRDREPORT_SESSION_LOCK = threading.RLock()
BIRDREPORT_OPENER = None

class BirdreportCaptchaRequired(RuntimeError):
    pass

def get_birdreport_opener():
    global BIRDREPORT_OPENER
    if BIRDREPORT_OPENER is None:
        ssl_context = ssl.create_default_context(cafile=certifi.where() if certifi else None)
        BIRDREPORT_OPENER = urllib.request.build_opener(
            urllib.request.HTTPCookieProcessor(BIRDREPORT_COOKIE_JAR),
            urllib.request.HTTPSHandler(context=ssl_context),
        )
    return BIRDREPORT_OPENER

def birdreport_request_bytes(req, timeout=20):
    with BIRDREPORT_SESSION_LOCK:
        with get_birdreport_opener().open(req, timeout=timeout) as response:
            return response.read(), response.headers.get("Content-Type")

def make_birdreport_payload(params):
    """模拟 birdreport 官网前端的 RSA 加密和请求签名。"""
    query = urllib.parse.urlencode(params)
    query_data = {}
    for item in query.split("&"):
        if "=" in item:
            key, value = item.split("=", 1)
            query_data[key] = value
        else:
            query_data[item] = ""

    plain = json.dumps(
        {key: query_data[key] for key in sorted(query_data)},
        ensure_ascii=False,
        separators=(",", ":")
    )
    public_key = serialization.load_der_public_key(base64.b64decode(BIRDREPORT_PUBLIC_KEY))
    encrypted = b"".join(
        public_key.encrypt(plain[index:index + 117].encode("utf-8"), asym_padding.PKCS1v15())
        for index in range(0, len(plain), 117)
    )
    timestamp = str(int(time.time() * 1000))
    request_id = uuid.uuid4().hex
    sign = hashlib.md5((plain + request_id + timestamp).encode("utf-8")).hexdigest()
    return base64.b64encode(encrypted), {
        "timestamp": timestamp,
        "requestId": request_id,
        "sign": sign,
        "User-Agent": "Mozilla/5.0",
        "Referer": "https://www.birdreport.cn/home/search/page.html",
        "Origin": "https://www.birdreport.cn",
        "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8",
    }

def decode_birdreport_data(data):
    encrypted = base64.b64decode(data)
    cipher = Cipher(algorithms.AES(BIRDREPORT_AES_KEY), modes.CBC(BIRDREPORT_AES_IV))
    decryptor = cipher.decryptor()
    padded = decryptor.update(encrypted) + decryptor.finalize()
    unpadder = sym_padding.PKCS7(128).unpadder()
    plain = unpadder.update(padded) + unpadder.finalize()
    return json.loads(plain.decode("utf-8"))

def birdreport_post(path, params):
    body, headers = make_birdreport_payload(params)
    req = urllib.request.Request(
        urllib.parse.urljoin(BIRDREPORT_API, path),
        data=body,
        headers=headers,
        method="POST"
    )
    response_body, _ = birdreport_request_bytes(req)
    return json.loads(response_body.decode("utf-8"))

def describe_birdreport_error(response, fallback):
    code = response.get("code")
    msg = response.get("msg")
    if code in (405, 505):
        return f"观鸟数据中心返回 code={code}，触发访问频率校验，请在弹窗输入验证码后继续查询"
    if msg:
        return f"观鸟数据中心返回 code={code}: {msg}"
    return f"观鸟数据中心返回 code={code}: {fallback}"

def raise_birdreport_error(response, fallback):
    message = describe_birdreport_error(response, fallback)
    if response.get("code") in (405, 505):
        raise BirdreportCaptchaRequired(message)
    raise RuntimeError(message)

def merge_bird_taxon_records(record_groups):
    merged = {}
    total_records = 0
    for outside_type, records in record_groups:
        for record in records:
            taxon_id = record.get("taxon_id") or record.get("taxonid") or record.get("taxonname")
            key = str(taxon_id)
            record_count = int(record.get("recordcount") or record.get("record_count") or 0)
            total_records += record_count
            if key not in merged:
                item = dict(record)
                item["outside_type"] = outside_type
                item["recordcount"] = record_count
                merged[key] = item
            else:
                merged[key]["recordcount"] = int(merged[key].get("recordcount") or 0) + record_count
                if outside_type == 1:
                    merged[key]["outside_type"] = 1
    records = list(merged.values())
    records.sort(key=lambda item: (item.get("outside_type", 0), item.get("taxon_id") or 0))
    return records, total_records

def init_bird_records_tables():
    """初始化鸟种记录相关表。"""
    if not _manager.db_enabled:
        return
    try:
        conn = get_db()
        with conn.cursor() as cur:
            cur.execute("""
                CREATE TABLE IF NOT EXISTS bird_locations (
                    id INT AUTO_INCREMENT PRIMARY KEY,
                    location_name VARCHAR(255) NOT NULL UNIQUE,
                    enabled TINYINT(1) NOT NULL DEFAULT 1,
                    sort_order INT NOT NULL DEFAULT 0,
                    created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
                    updated_at DATETIME DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP
                ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
            """)
            cur.execute("""
                CREATE TABLE IF NOT EXISTS bird_record_summary (
                    id BIGINT AUTO_INCREMENT PRIMARY KEY,
                    batch_no VARCHAR(32) NOT NULL,
                    location_name VARCHAR(255) NOT NULL,
                    start_date DATE NOT NULL,
                    end_date DATE NOT NULL,
                    species_count INT NOT NULL DEFAULT 0,
                    report_count INT NOT NULL DEFAULT 0,
                    record_count INT NOT NULL DEFAULT 0,
                    individual_count INT DEFAULT NULL,
                    query_time DATETIME NOT NULL,
                    created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
                    INDEX idx_batch_species (batch_no, species_count),
                    INDEX idx_location_time (location_name, query_time)
                ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
            """)
            cur.execute("""
                CREATE TABLE IF NOT EXISTS bird_record_detail (
                    id BIGINT AUTO_INCREMENT PRIMARY KEY,
                    summary_id BIGINT NOT NULL,
                    batch_no VARCHAR(32) NOT NULL,
                    report_no VARCHAR(100) NOT NULL,
                    observation_location VARCHAR(500) NOT NULL,
                    bird_name VARCHAR(255) NOT NULL,
                    protection_level VARCHAR(20) DEFAULT NULL,
                    bird_count INT DEFAULT NULL,
                    record_user VARCHAR(255) NOT NULL,
                    observation_time VARCHAR(100) DEFAULT NULL,
                    outside_type TINYINT DEFAULT 0,
                    source_state TINYINT DEFAULT NULL,
                    created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
                    INDEX idx_summary (summary_id),
                    INDEX idx_batch_location (batch_no, observation_location),
                    INDEX idx_bird_name (bird_name)
                ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
            """)
            cur.execute("""
                CREATE TABLE IF NOT EXISTS bird_record_source_detail (
                    id BIGINT AUTO_INCREMENT PRIMARY KEY,
                    batch_no VARCHAR(32) NOT NULL,
                    source_region VARCHAR(255) NOT NULL,
                    start_date DATE NOT NULL,
                    end_date DATE NOT NULL,
                    report_no VARCHAR(100) NOT NULL,
                    observation_location VARCHAR(500) NOT NULL,
                    bird_name VARCHAR(255) NOT NULL,
                    protection_level VARCHAR(20) DEFAULT NULL,
                    bird_count INT DEFAULT NULL,
                    record_user VARCHAR(255) NOT NULL,
                    observation_time VARCHAR(100) DEFAULT NULL,
                    outside_type TINYINT DEFAULT 0,
                    source_state TINYINT DEFAULT NULL,
                    query_time DATETIME NOT NULL,
                    created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
                    INDEX idx_source_batch_location (batch_no, observation_location),
                    INDEX idx_source_bird_name (bird_name),
                    INDEX idx_source_query_time (query_time)
                ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
            """)
            for table_name in ("bird_record_detail", "bird_record_source_detail"):
                try:
                    cur.execute(f"ALTER TABLE {table_name} ADD COLUMN protection_level VARCHAR(20) DEFAULT NULL")
                except Exception:
                    pass
            try:
                cur.execute(
                    "ALTER TABLE bird_record_detail DROP FOREIGN KEY fk_bird_record_detail_summary"
                )
            except Exception:
                pass
            try:
                cur.execute("""
                    UPDATE bird_record_detail d
                    INNER JOIN protected_wildlife_catalog p ON p.chinese_name = d.bird_name
                    SET d.protection_level = p.protection_level
                    WHERE d.protection_level IS NULL
                """)
                cur.execute("""
                    UPDATE bird_record_source_detail d
                    INNER JOIN protected_wildlife_catalog p ON p.chinese_name = d.bird_name
                    SET d.protection_level = p.protection_level
                    WHERE d.protection_level IS NULL
                """)
            except Exception:
                pass
            cur.execute("SELECT COUNT(*) FROM bird_locations")
            if int(cur.fetchone()[0] or 0) == 0:
                cur.execute(
                    "INSERT INTO bird_locations (location_name, sort_order) VALUES (%s, %s)",
                    ("奥森", 1)
                )
        conn.commit()
        conn.close()
    except Exception as e:
        print(f"[DB] 初始化鸟种记录表失败: {e}")

def dict_rows(cursor):
    columns = [desc[0] for desc in cursor.description]
    return [dict(zip(columns, row)) for row in cursor.fetchall()]

def parse_bird_count(value):
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    match = re.search(r"\d+", text)
    if not match:
        return None
    return int(match.group(0))

def get_bird_record_blacklist():
    """读取观鸟记录用户黑名单，并统一去重、忽略大小写匹配。"""
    raw = _manager.get_all().get("bird_record_blacklist", [])
    if isinstance(raw, str):
        raw = re.split(r"[,，\n]", raw)
    result = []
    seen = set()
    for value in raw or []:
        name = str(value or "").strip()
        key = name.casefold()
        if name and key not in seen:
            seen.add(key)
            result.append(name)
    return result

def filter_bird_record_blacklist(details):
    blocked = {name.casefold() for name in get_bird_record_blacklist()}
    if not blocked:
        return details
    return [item for item in details
            if str(item.get("record_user") or "").strip().casefold() not in blocked]

def join_bird_location(record, fallback_location):
    province = str(record.get("province_name") or "").strip()
    city = str(record.get("city_name") or "").strip()
    district = str(record.get("district_name") or "").strip()
    point_name = str(record.get("point_name") or record.get("pointname") or fallback_location or "").strip()
    parts = [province]
    if city and city != province:
        parts.append(city)
    parts.extend([district, point_name])
    return "".join(str(part).strip() for part in parts if part)

def normalize_bird_detail(record, fallback_location):
    state = record.get("state")
    if str(state) != "2":
        return None
    point_name = str(record.get("point_name") or record.get("pointname") or "").strip()
    if re.fullmatch(r"\d+", point_name):
        return None
    report_no = str(record.get("serial_id") or "").strip()
    bird_name = str(record.get("taxon_name") or record.get("taxonname") or "").strip()
    record_user = str(record.get("username") or "").strip()
    if not report_no or not bird_name:
        return None
    uncertainty_value = record.get("taxon_uncertain")
    if uncertainty_value is None:
        uncertainty_value = record.get("uncertain", record.get("is_uncertain"))
    uncertain_text = bird_name.casefold()
    bird_uncertain = bool(uncertainty_value) or any(marker in uncertain_text for marker in ("?", "？", "疑似", "待定", "sp.", " cf.", " cf"))
    start_time = str(record.get("start_time") or "").strip()
    end_time = str(record.get("end_time") or "").strip()
    if start_time and end_time:
        observation_time = f"{start_time[:16]} 至 {end_time[:16]}"
    else:
        observation_time = start_time[:16] or None
    return {
        "report_no": report_no,
        "observation_location": join_bird_location(record, fallback_location),
        "bird_name": bird_name,
        "bird_uncertain": bird_uncertain,
        "bird_count": parse_bird_count(record.get("taxon_count") or record.get("taxoncount")),
        "record_user": record_user,
        "observation_time": observation_time,
        "outside_type": int(record.get("outside_type") or 0),
        "source_state": int(state),
    }

def annotate_bird_protection_levels(conn, details):
    """按国家重点保护野生动物名录给查询明细补充保护级别。"""
    names = sorted({item.get("bird_name") for item in details if item.get("bird_name")})
    if not names:
        return
    if not _manager.db_enabled:
        levels = get_protected_wildlife_level(names)
        for item in details:
            item["protection_level"] = levels.get(item.get("bird_name"))
        return
    placeholders = ", ".join(["%s"] * len(names))
    with conn.cursor() as cur:
        cur.execute(
            f"SELECT chinese_name, protection_level FROM protected_wildlife_catalog "
            f"WHERE chinese_name IN ({placeholders})",
            names,
        )
        levels = {row[0]: row[1] for row in cur.fetchall()}
    for item in details:
        item["protection_level"] = levels.get(item.get("bird_name"))

def fetch_birdreport_public_details(location, start_date, end_date):
    details = []
    for outside_type in (0, 1):
        page = 1
        limit = 1500
        while True:
            params = {
                "startTime": start_date,
                "endTime": end_date,
                "pointname": location,
                "state": "2",
                "version": "CH4",
                "mode": "0",
                "outside_type": str(outside_type),
                "page": str(page),
                "limit": str(limit),
            }
            response = birdreport_post("front/record/search/page", params)
            if response.get("code") not in (0, "0"):
                raise_birdreport_error(response, "birdreport 明细查询失败")
            total_count = int(response.get("count") or 0)
            if total_count <= 0:
                break
            records = decode_birdreport_data(response.get("data"))
            for record in records:
                record["outside_type"] = outside_type
                detail = normalize_bird_detail(record, location)
                if detail:
                    details.append(detail)
            if page * limit >= total_count or not records:
                break
            page += 1
    return details

def build_bird_region_name(province, city=""):
    province = (province or "").strip()
    city = (city or "").strip()
    return f"{province}{city}" if city else province

def fetch_birdreport_public_area_details(province, start_date, end_date, city=""):
    region_name = build_bird_region_name(province, city)
    details = []
    for outside_type in (0, 1):
        page = 1
        limit = 1500
        while True:
            params = {
                "startTime": start_date,
                "endTime": end_date,
                "province": province,
                "state": "2",
                "version": "CH4",
                "mode": "0",
                "outside_type": str(outside_type),
                "page": str(page),
                "limit": str(limit),
            }
            if city:
                params["city"] = city
            response = birdreport_post("front/record/search/page", params)
            if response.get("code") not in (0, "0"):
                raise_birdreport_error(response, "birdreport 区域明细查询失败")
            total_count = int(response.get("count") or 0)
            if total_count <= 0:
                break
            records = decode_birdreport_data(response.get("data"))
            for record in records:
                record["outside_type"] = outside_type
                detail = normalize_bird_detail(record, region_name)
                if detail:
                    details.append(detail)
            if page * limit >= total_count or not records:
                break
            page += 1
    return details

def group_bird_details_by_location(details):
    grouped = {}
    for detail in details:
        location = detail.get("observation_location") or ""
        if not location:
            continue
        grouped.setdefault(location, []).append(detail)
    ranked = []
    for location, location_details in grouped.items():
        species_count = len({item["bird_name"] for item in location_details if item.get("bird_name")})
        report_count = len({item["report_no"] for item in location_details if item.get("report_no")})
        ranked.append({
            "location": location,
            "details": location_details,
            "species_count": species_count,
            "report_count": report_count,
            "record_count": len(location_details),
        })
    ranked.sort(key=lambda item: (item["species_count"], item["record_count"], item["report_count"]), reverse=True)
    return ranked

def sync_bird_locations(conn, ranked_locations):
    with conn.cursor() as cur:
        cur.execute("UPDATE bird_locations SET enabled = 0")
        for index, item in enumerate(ranked_locations, start=1):
            cur.execute(
                """
                INSERT INTO bird_locations (location_name, enabled, sort_order)
                VALUES (%s, 1, %s)
                ON DUPLICATE KEY UPDATE
                    enabled = 1,
                    sort_order = VALUES(sort_order),
                    updated_at = CURRENT_TIMESTAMP
                """,
                (item["location"], index)
            )

def generate_bird_batch_no(conn, query_time):
    prefix = query_time.strftime("%Y%m%d%H%M%S")
    with conn.cursor() as cur:
        cur.execute(
            "SELECT MAX(batch_no) FROM bird_record_summary WHERE batch_no LIKE %s",
            (prefix + "%",)
        )
        row = cur.fetchone()
    last_batch = row[0] if row else None
    next_seq = int(str(last_batch)[-3:]) + 1 if last_batch else 1
    return f"{prefix}{next_seq:03d}"


def generate_memory_batch_no():
    return "MEMORY-" + datetime.now().strftime("%Y%m%d%H%M%S")

def save_bird_source_details(conn, batch_no, source_region, start_date, end_date, query_time, details):
    if not details:
        return 0
    with conn.cursor() as cur:
        cur.executemany(
            """
            INSERT INTO bird_record_source_detail (
                batch_no, source_region, start_date, end_date, report_no,
                observation_location, bird_name, protection_level, bird_count, record_user,
                observation_time, outside_type, source_state, query_time
            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            """,
            [
                (
                    batch_no,
                    source_region,
                    start_date,
                    end_date,
                    item["report_no"],
                    item["observation_location"],
                    item["bird_name"],
                    item.get("protection_level"),
                    item["bird_count"],
                    item["record_user"],
                    item["observation_time"],
                    item["outside_type"],
                    item["source_state"],
                    query_time.strftime("%Y-%m-%d %H:%M:%S"),
                )
                for item in details
            ]
        )
    return len(details)

def save_bird_location_summary(conn, batch_no, location, start_date, end_date, query_time, details):
    species = {item["bird_name"] for item in details if item.get("bird_name")}
    reports = {item["report_no"] for item in details if item.get("report_no")}
    counted_values = [item["bird_count"] for item in details if item.get("bird_count") is not None]
    individual_count = sum(counted_values) if counted_values else None
    with conn.cursor() as cur:
        cur.execute(
            """
            INSERT INTO bird_record_summary (
                batch_no, location_name, start_date, end_date, species_count,
                report_count, record_count, individual_count, query_time
            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
            """,
            (
                batch_no,
                location,
                start_date,
                end_date,
                len(species),
                len(reports),
                len(details),
                individual_count,
                query_time.strftime("%Y-%m-%d %H:%M:%S"),
            )
        )
        summary_id = cur.lastrowid
        cur.executemany(
            """
            INSERT INTO bird_record_detail (
                summary_id, batch_no, report_no, observation_location, bird_name,
                protection_level, bird_count, record_user, observation_time, outside_type, source_state
            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            """,
            [
                (
                    summary_id,
                    batch_no,
                    item["report_no"],
                    item["observation_location"],
                    item["bird_name"],
                    item.get("protection_level"),
                    item["bird_count"],
                    item["record_user"],
                    item["observation_time"],
                    item["outside_type"],
                    item["source_state"],
                )
                for item in details
            ]
        )
    return {
        "id": summary_id,
        "batch_no": batch_no,
        "location_name": location,
        "species_count": len(species),
        "report_count": len(reports),
        "record_count": len(details),
        "individual_count": individual_count,
    }

def validate_bird_dates(start_date, end_date):
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", start_date or ""):
        raise ValueError("起始日期格式必须为 YYYY-MM-DD")
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", end_date or ""):
        raise ValueError("结束日期格式必须为 YYYY-MM-DD")
    if start_date > end_date:
        raise ValueError("起始日期不能晚于结束日期")

def excel_column_name(index):
    name = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        name = chr(65 + remainder) + name
    return name

def excel_safe_sheet_name(name, fallback):
    cleaned = re.sub(r"[\[\]\*\?/\\:]", "", str(name or "")).strip()
    return (cleaned or fallback)[:31]

def excel_cell_xml(row_index, col_index, value, header=False):
    ref = f"{excel_column_name(col_index)}{row_index}"
    style = ' s="1"' if header else ""
    if value is None:
        return f'<c r="{ref}"{style}/>'
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return f'<c r="{ref}"{style}><v>{value}</v></c>'
    text = xml_escape(str(value), {'"': "&quot;", "'": "&apos;"})
    return f'<c r="{ref}"{style} t="inlineStr"><is><t>{text}</t></is></c>'

def excel_display_width(value):
    """按 Excel 中英文字符的大致显示宽度估算列宽。"""
    return sum(2 if unicodedata.east_asian_width(char) in ("W", "F", "A") else 1 for char in str(value))

def build_excel_sheet_xml(rows):
    max_cols = max((len(row) for row in rows), default=1)
    max_rows = max(len(rows), 1)
    dimension = f"A1:{excel_column_name(max_cols)}{max_rows}"
    widths = []
    for col_index in range(1, max_cols + 1):
        max_width = 10
        for row in rows[:200]:
            if col_index <= len(row) and row[col_index - 1] is not None:
                max_width = max(max_width, min(excel_display_width(row[col_index - 1]) + 2, 50))
        widths.append(f'<col min="{col_index}" max="{col_index}" width="{max_width}" customWidth="1"/>')
    row_xml = []
    for row_index, row in enumerate(rows, start=1):
        cells = "".join(
            excel_cell_xml(row_index, col_index, value, header=row_index == 1)
            for col_index, value in enumerate(row, start=1)
        )
        height = ' ht="24" customHeight="1"' if row_index == 1 else ""
        row_xml.append(f'<row r="{row_index}"{height}>{cells}</row>')
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f'<dimension ref="{dimension}"/>'
        '<sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/>'
        '<selection pane="bottomLeft" activeCell="A2" sqref="A2"/></sheetView></sheetViews>'
        f'<cols>{"".join(widths)}</cols>'
        f'<sheetData>{"".join(row_xml)}</sheetData>'
        '</worksheet>'
    )

def build_bird_records_xlsx(sheets):
    buffer = BytesIO()
    sheet_names = [excel_safe_sheet_name(name, f"Sheet{index}") for index, (name, _) in enumerate(sheets, start=1)]
    now = datetime.now().replace(microsecond=0).isoformat()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as workbook:
        sheet_overrides = "".join(
            f'<Override PartName="/xl/worksheets/sheet{index}.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            for index in range(1, len(sheets) + 1)
        )
        workbook.writestr("[Content_Types].xml", (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
            '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>'
            '<Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>'
            f'{sheet_overrides}</Types>'
        ))
        workbook.writestr("_rels/.rels", (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
            '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>'
            '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties" Target="docProps/app.xml"/>'
            '</Relationships>'
        ))
        workbook.writestr("docProps/core.xml", (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
            'xmlns:dc="http://purl.org/dc/elements/1.1/" '
            'xmlns:dcterms="http://purl.org/dc/terms/" '
            'xmlns:dcmitype="http://purl.org/dc/dcmitype/" '
            'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
            '<dc:creator>ToolBox</dc:creator>'
            '<cp:lastModifiedBy>ToolBox</cp:lastModifiedBy>'
            f'<dcterms:created xsi:type="dcterms:W3CDTF">{now}</dcterms:created>'
            f'<dcterms:modified xsi:type="dcterms:W3CDTF">{now}</dcterms:modified>'
            '</cp:coreProperties>'
        ))
        workbook.writestr("docProps/app.xml", (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
            'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
            '<Application>ToolBox</Application>'
            f'<TitlesOfParts><vt:vector size="{len(sheet_names)}" baseType="lpstr">'
            + ''.join(f'<vt:lpstr>{xml_escape(name)}</vt:lpstr>' for name in sheet_names) +
            '</vt:vector></TitlesOfParts>'
            '</Properties>'
        ))
        workbook.writestr("xl/styles.xml", (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<fonts count="2"><font><sz val="11"/><name val="Calibri"/></font>'
            '<font><b/><color rgb="FFFFFFFF"/><sz val="11"/><name val="Calibri"/></font></fonts>'
            '<fills count="3"><fill><patternFill patternType="none"/></fill>'
            '<fill><patternFill patternType="gray125"/></fill>'
            '<fill><patternFill patternType="solid"><fgColor rgb="FF17705E"/><bgColor indexed="64"/></patternFill></fill></fills>'
            '<borders count="1"><border/></borders>'
            '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
            '<cellXfs count="2"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
            '<xf numFmtId="0" fontId="1" fillId="2" borderId="0" xfId="0" applyFont="1" applyFill="1"/></cellXfs>'
            '<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>'
            '</styleSheet>'
        ))
        workbook_sheets = "".join(
            '<sheet name="{name}" sheetId="{index}" r:id="rId{index}"/>'.format(
                name=xml_escape(name, {'"': "&quot;"}),
                index=index,
            )
            for index, name in enumerate(sheet_names, start=1)
        )
        workbook.writestr("xl/workbook.xml", (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            f'<sheets>{workbook_sheets}</sheets>'
            '</workbook>'
        ))
        workbook_rels = "".join(
            f'<Relationship Id="rId{index}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
            f'Target="worksheets/sheet{index}.xml"/>'
            for index in range(1, len(sheets) + 1)
        )
        workbook_rels += (
            f'<Relationship Id="rId{len(sheets) + 1}" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
        )
        workbook.writestr("xl/_rels/workbook.xml.rels", (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'{workbook_rels}</Relationships>'
        ))
        for index, (_, rows) in enumerate(sheets, start=1):
            workbook.writestr(f"xl/worksheets/sheet{index}.xml", build_excel_sheet_xml(rows))
    buffer.seek(0)
    return buffer

def get_latest_bird_batch_export_data(conn):
    with conn.cursor() as cur:
        cur.execute("SELECT MAX(batch_no) FROM bird_record_source_detail")
        row = cur.fetchone()
        batch_no = row[0] if row else None
        source_table = "bird_record_source_detail"
        if not batch_no:
            cur.execute("SELECT MAX(batch_no) FROM bird_record_detail")
            row = cur.fetchone()
            batch_no = row[0] if row else None
            source_table = "bird_record_detail"
            if not batch_no:
                return None
        cur.execute("""
            SELECT d.observation_location, d.bird_name,
                   d.protection_level, p.protection_level AS catalog_protection_level,
                   d.bird_count, d.record_user, d.report_no, d.observation_time,
                   d.outside_type, d.source_state
            FROM {source_table} d
            LEFT JOIN protected_wildlife_catalog p ON p.chinese_name = d.bird_name
            WHERE d.batch_no = %s
            ORDER BY d.observation_location ASC, d.bird_name ASC, d.report_no ASC, d.id ASC
        """.format(source_table=source_table), (batch_no,))
        details = filter_bird_record_blacklist(dict_rows(cur))
        # 黑名单变化后，导出中的地点统计也按过滤后的明细重新计算。
        grouped = {}
        for detail in details:
            grouped.setdefault(detail.get("observation_location") or "", []).append(detail)
        all_location_summaries = [
            {
                "location_name": location,
                "species_count": len({item.get("bird_name") for item in rows if item.get("bird_name")}),
                "report_count": len({item.get("report_no") for item in rows if item.get("report_no")}),
                "record_count": len(rows),
                "individual_count": sum(item.get("bird_count") or 0 for item in rows),
            }
            for location, rows in grouped.items() if location
        ]
        all_location_summaries.sort(key=lambda item: (-item["species_count"], -item["record_count"], item["location_name"]))
        cur.execute("""
            SELECT location_name, species_count, report_count, record_count,
                   individual_count, start_date, end_date, query_time
            FROM bird_record_summary
            WHERE batch_no = %s
            ORDER BY species_count DESC, record_count DESC, id ASC
        """, (batch_no,))
        summaries = dict_rows(cur)
        cur.execute("""
            SELECT d.observation_location, d.bird_name,
                   COALESCE(d.protection_level, p.protection_level) AS protection_level,
                   COUNT(*) AS record_count,
                   COUNT(DISTINCT report_no) AS report_count,
                   SUM(CASE WHEN bird_count IS NULL THEN 0 ELSE bird_count END) AS individual_count
            FROM {source_table} d
            LEFT JOIN protected_wildlife_catalog p ON p.chinese_name = d.bird_name
            WHERE d.batch_no = %s
              AND COALESCE(d.protection_level, p.protection_level) IN ('Ⅰ级', 'Ⅱ级')
            GROUP BY d.observation_location, d.bird_name,
                     COALESCE(d.protection_level, p.protection_level)
            ORDER BY d.observation_location ASC, d.bird_name ASC
        """.format(source_table=source_table), (batch_no,))
        bird_summaries = dict_rows(cur)
        protected_groups = {}
        for detail in details:
            level = detail.get("protection_level") or detail.get("catalog_protection_level")
            if level in ("Ⅰ级", "Ⅱ级") and detail.get("bird_name"):
                protected_groups.setdefault((detail.get("observation_location"), detail.get("bird_name"), level), []).append(detail)
        bird_summaries = [
            {
                "observation_location": key[0], "bird_name": key[1], "protection_level": key[2],
                "record_count": len(rows),
                "report_count": len({item.get("report_no") for item in rows if item.get("report_no")}),
                "individual_count": sum(item.get("bird_count") or 0 for item in rows),
            }
            for key, rows in sorted(protected_groups.items())
        ]
        summary_stats = {row["location_name"]: row for row in all_location_summaries}
        summaries = [
            dict(row, **summary_stats[row["location_name"]])
            for row in summaries if row.get("location_name") in summary_stats
        ]
    return {
        "batch_no": batch_no,
        "source_table": source_table,
        "details": details,
        "all_location_summaries": all_location_summaries,
        "summaries": summaries,
        "bird_summaries": bird_summaries,
    }

def get_memory_bird_batch_export_data(memory_batch):
    """把无数据库模式下的最近一次内存查询整理成与数据库导出相同的结构。"""
    if not memory_batch or memory_batch.get("expires_at", 0) <= time.monotonic():
        return None
    result = memory_batch.get("result") or {}
    batch_no = result.get("batchNo")
    details = filter_bird_record_blacklist(list(memory_batch.get("details") or []))
    if not batch_no:
        return None

    grouped_locations = {}
    protected_groups = {}
    for detail in details:
        location = detail.get("observation_location") or ""
        if not location:
            continue
        grouped_locations.setdefault(location, []).append(detail)
        protection_level = detail.get("protection_level")
        bird_name = detail.get("bird_name") or ""
        if protection_level in ("Ⅰ级", "Ⅱ级") and bird_name:
            protected_groups.setdefault((location, bird_name, protection_level), []).append(detail)

    all_location_summaries = []
    for location, location_details in grouped_locations.items():
        all_location_summaries.append({
            "location_name": location,
            "species_count": len({item.get("bird_name") for item in location_details if item.get("bird_name")}),
            "report_count": len({item.get("report_no") for item in location_details if item.get("report_no")}),
            "record_count": len(location_details),
            "individual_count": sum(item.get("bird_count") or 0 for item in location_details),
        })
    all_location_summaries.sort(
        key=lambda item: (-item["species_count"], -item["record_count"], item["location_name"])
    )

    bird_summaries = []
    for (location, bird_name, protection_level), bird_details in sorted(protected_groups.items()):
        bird_summaries.append({
            "observation_location": location,
            "bird_name": bird_name,
            "protection_level": protection_level,
            "record_count": len(bird_details),
            "report_count": len({item.get("report_no") for item in bird_details if item.get("report_no")}),
            "individual_count": sum(item.get("bird_count") or 0 for item in bird_details),
        })

    return {
        "batch_no": batch_no,
        "source_table": "memory",
        "details": details,
        "all_location_summaries": all_location_summaries,
        "summaries": list(result.get("saved") or []),
        "bird_summaries": bird_summaries,
    }

def format_excel_value(value):
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value

def build_bird_records_export_workbook(export_data):
    batch_no = export_data["batch_no"]
    source_url = "https://www.birdreport.cn/home/search/page.html"
    detail_rows = [[
        "观测地点", "鸟名", "保护级别", "数量", "记录用户", "报告编号", "观测时间", "来源类型", "公开状态", "批次号", "数据来源"
    ]]
    for row in export_data["details"]:
        protection_level = row.get("protection_level") or row.get("catalog_protection_level")
        detail_rows.append([
            row.get("observation_location"),
            row.get("bird_name"),
            protection_level,
            row.get("bird_count"),
            row.get("record_user"),
            row.get("report_no"),
            row.get("observation_time"),
            row.get("outside_type"),
            row.get("source_state"),
            batch_no,
            source_url,
        ])
    all_location_rows = [[
        "排名", "观测地点", "鸟种数量", "报告数量", "明细数量", "个体数量", "批次号"
    ]]
    for index, row in enumerate(export_data["all_location_summaries"], start=1):
        all_location_rows.append([
            index,
            row.get("location_name"),
            row.get("species_count"),
            row.get("report_count"),
            row.get("record_count"),
            row.get("individual_count"),
            batch_no,
        ])
    top_location_rows = [[
        "排名", "观测地点", "鸟种数量", "报告数量", "明细数量", "个体数量", "起始日期", "结束日期", "查询时间", "批次号"
    ]]
    for index, row in enumerate(export_data["summaries"], start=1):
        top_location_rows.append([
            index,
            row.get("location_name"),
            row.get("species_count"),
            row.get("report_count"),
            row.get("record_count"),
            row.get("individual_count"),
            format_excel_value(row.get("start_date")),
            format_excel_value(row.get("end_date")),
            format_excel_value(row.get("query_time")),
            batch_no,
        ])
    bird_rows = [["观测地点", "鸟名", "保护级别", "记录条数", "报告数量", "个体数量", "批次号"]]
    for row in export_data["bird_summaries"]:
        bird_rows.append([
            row.get("observation_location"),
            row.get("bird_name"),
            row.get("protection_level"),
            row.get("record_count"),
            row.get("report_count"),
            row.get("individual_count"),
            batch_no,
        ])
    return build_bird_records_xlsx([
        ("数据来源明细", detail_rows),
        ("全量地点汇总", all_location_rows),
        ("Top20地点汇总", top_location_rows),
        ("鸟种汇总", bird_rows),
    ])

@app.route('/api/bird-records/blacklist', methods=['GET'])
def api_bird_records_blacklist():
    return jsonify({"users": get_bird_record_blacklist()})

@app.route('/api/bird-records/blacklist', methods=['PUT'])
def api_bird_records_blacklist_update():
    data = request.get_json(silent=True) or {}
    users = data.get("users", data.get("user", []))
    if isinstance(users, str):
        users = re.split(r"[,，\n]", users)
    normalized = []
    seen = set()
    for value in users or []:
        name = str(value or "").strip()
        key = name.casefold()
        if name and key not in seen:
            seen.add(key)
            normalized.append(name)
    _manager.update({"bird_record_blacklist": normalized})
    with _CACHE_LOCK:
        for key in list(_CACHE):
            if key.startswith(_BIRD_QUERY_CACHE_PREFIX):
                _CACHE.pop(key, None)
    return jsonify({"users": normalized, "message": "黑名单已更新，后续展示和 AI 查询将自动排除这些用户"})

@app.route('/api/bird-records/locations', methods=['GET'])
def api_bird_locations():
    init_bird_records_tables()
    try:
        conn = get_db()
        with conn.cursor() as cur:
            cur.execute("""
                SELECT id, location_name, sort_order, created_at
                FROM bird_locations
                WHERE enabled = 1
                ORDER BY sort_order ASC, id ASC
            """)
            rows = dict_rows(cur)
        conn.close()
        return jsonify(rows)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/api/bird-records/locations', methods=['POST'])
def api_bird_locations_add():
    init_bird_records_tables()
    data = request.get_json(silent=True) or {}
    location_name = str(data.get("location_name") or data.get("location") or "").strip()
    if not location_name:
        return jsonify({"error": "地点不能为空"}), 400
    try:
        conn = get_db()
        with conn.cursor() as cur:
            cur.execute("SELECT COALESCE(MAX(sort_order), 0) + 1 FROM bird_locations")
            sort_order = int(cur.fetchone()[0] or 1)
            cur.execute(
                """
                INSERT INTO bird_locations (location_name, enabled, sort_order)
                VALUES (%s, 1, %s)
                ON DUPLICATE KEY UPDATE enabled = 1, updated_at = CURRENT_TIMESTAMP
                """,
                (location_name, sort_order)
            )
        conn.commit()
        conn.close()
        return api_bird_locations()
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/api/bird-records/locations/<int:location_id>', methods=['DELETE'])
def api_bird_locations_delete(location_id):
    init_bird_records_tables()
    try:
        conn = get_db()
        with conn.cursor() as cur:
            cur.execute("UPDATE bird_locations SET enabled = 0 WHERE id = %s", (location_id,))
        conn.commit()
        conn.close()
        return jsonify({"status": "ok"})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/api/bird-records/captcha', methods=['GET'])
def api_bird_records_captcha():
    try:
        timestamp = str(int(time.time() * 1000))
        req = urllib.request.Request(
            urllib.parse.urljoin(BIRDREPORT_API, f"front/code/visited/generate?timestamp={timestamp}"),
            headers={
                "User-Agent": "Mozilla/5.0",
                "Referer": "https://www.birdreport.cn/home/code/verify.html",
                "Origin": "https://www.birdreport.cn",
            },
        )
        image, content_type = birdreport_request_bytes(req)
        response = send_file(
            BytesIO(image),
            mimetype=content_type or "image/jpeg",
            max_age=0,
        )
        response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        return response
    except Exception as e:
        return jsonify({"error": f"获取验证码失败: {e}"}), 502

@app.route('/api/bird-records/captcha/verify', methods=['POST'])
def api_bird_records_captcha_verify():
    data = request.get_json(silent=True) or {}
    code = str(data.get("code") or "").strip()
    if not re.fullmatch(r"[A-Za-z0-9]{4,6}", code):
        return jsonify({"error": "请输入正确的验证码"}), 400
    try:
        body = json.dumps({"code": code}, ensure_ascii=False).encode("utf-8")
        req = urllib.request.Request(
            urllib.parse.urljoin(BIRDREPORT_API, "front/code/visited/verify"),
            data=body,
            headers={
                "User-Agent": "Mozilla/5.0",
                "Referer": "https://www.birdreport.cn/home/code/verify.html",
                "Origin": "https://www.birdreport.cn",
                "Content-Type": "application/json",
            },
            method="POST",
        )
        response_body, _ = birdreport_request_bytes(req)
        result = json.loads(response_body.decode("utf-8"))
        if result.get("success"):
            return jsonify({"status": "ok", "message": result.get("msg") or "验证码已通过"})
        return jsonify({"error": result.get("msg") or "验证码验证失败"}), 400
    except Exception as e:
        return jsonify({"error": f"验证码验证失败: {e}"}), 502

@app.route('/api/bird-records/query-batch', methods=['POST'])
def api_bird_records_query_batch():
    init_bird_records_tables()
    today = datetime.now()
    default_start = (today - timedelta(days=1)).strftime("%Y-%m-%d")
    data = request.get_json(silent=True) or {}
    start_date = str(data.get("startDate") or default_start).strip()
    end_date = str(data.get("endDate") or today.strftime("%Y-%m-%d")).strip()
    province = str(data.get("province") or "北京市").strip()
    city = str(data.get("city") or "").strip()
    region_name = build_bird_region_name(province, city)
    try:
        validate_bird_dates(start_date, end_date)
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    if not province:
        return jsonify({"error": "查询区域不能为空"}), 400

    conn = None
    try:
        if not _manager.db_enabled:
            global _LATEST_MEMORY_BATCH
            all_details = filter_bird_record_blacklist(fetch_birdreport_public_area_details(province, start_date, end_date, city))
            annotate_bird_protection_levels(None, all_details)
            ranked_locations = group_bird_details_by_location(all_details)
            batch_no = generate_memory_batch_no()
            saved = []
            memory_summaries = {}
            for index, item in enumerate(ranked_locations[:20], start=1):
                item_details = item["details"]
                summary = {
                    "id": index,
                    "batch_no": batch_no,
                    "location_name": item["location"],
                    "start_date": start_date,
                    "end_date": end_date,
                    "species_count": len({d.get("bird_name") for d in item_details}),
                    "report_count": len({d.get("report_no") for d in item_details if d.get("report_no")}),
                    "record_count": len(item_details),
                    "individual_count": sum(d.get("bird_count") or 0 for d in item_details),
                    "query_time": datetime.now().isoformat(),
                }
                saved.append(summary)
                memory_summaries[index] = {
                    "summary": summary,
                    "details": item_details,
                    "birdNames": sorted({d.get("bird_name") for d in item_details if d.get("bird_name")}),
                }
            result = {
                "batchNo": batch_no,
                "province": province,
                "city": city,
                "regionName": region_name,
                "startDate": start_date,
                "endDate": end_date,
                "totalLocations": len(ranked_locations),
                "savedLocations": min(len(ranked_locations), 20),
                "recordTotal": len(all_details),
                "sourceSaved": len(all_details),
                "saved": saved,
            }
            _LATEST_MEMORY_BATCH = {"result": result, "summaries": memory_summaries,
                                    "details": all_details, "expires_at": time.monotonic() + CACHE_TTL_SECONDS}
            return jsonify(_cache_set(_query_cache_key(_BIRD_QUERY_CACHE_PREFIX, {
                "batch": province, "city": city, "start_date": start_date, "end_date": end_date,
            }), result))
        conn = get_db()
        query_time = datetime.now()
        batch_no = generate_bird_batch_no(conn, query_time)
        all_details = filter_bird_record_blacklist(fetch_birdreport_public_area_details(province, start_date, end_date, city))
        annotate_bird_protection_levels(conn, all_details)
        source_saved = save_bird_source_details(conn, batch_no, region_name, start_date, end_date, query_time, all_details)
        ranked_locations = group_bird_details_by_location(all_details)
        top_locations = ranked_locations[:20]
        sync_bird_locations(conn, top_locations)
        saved = []
        for item in top_locations:
            saved.append(save_bird_location_summary(
                conn,
                batch_no,
                item["location"],
                start_date,
                end_date,
                query_time,
                item["details"],
            ))
        conn.commit()
        conn.close()
        return jsonify({
            "batchNo": batch_no,
            "province": province,
            "city": city,
            "regionName": region_name,
            "startDate": start_date,
            "endDate": end_date,
            "totalLocations": len(ranked_locations),
            "savedLocations": len(saved),
            "recordTotal": len(all_details),
            "sourceSaved": source_saved,
            "saved": saved,
        })
    except BirdreportCaptchaRequired as e:
        try:
            if conn:
                conn.rollback()
                conn.close()
        except Exception:
            pass
        return jsonify({"error": str(e), "captchaRequired": True}), 429
    except Exception as e:
        try:
            if conn:
                conn.rollback()
                conn.close()
        except Exception:
            pass
        return jsonify({"error": str(e)}), 500

@app.route('/api/bird-records/latest', methods=['GET'])
def api_bird_records_latest():
    init_bird_records_tables()
    if not _manager.db_enabled:
        if _LATEST_MEMORY_BATCH and _LATEST_MEMORY_BATCH["expires_at"] > time.monotonic():
            result = _LATEST_MEMORY_BATCH["result"]
            top20 = []
            for entry in (_LATEST_MEMORY_BATCH.get("summaries") or {}).values():
                details = filter_bird_record_blacklist(entry.get("details") or [])
                if not details:
                    continue
                summary = dict(entry.get("summary") or {})
                summary.update({
                    "species_count": len({item.get("bird_name") for item in details if item.get("bird_name")}),
                    "report_count": len({item.get("report_no") for item in details if item.get("report_no")}),
                    "record_count": len(details),
                    "individual_count": sum(item.get("bird_count") or 0 for item in details),
                })
                top20.append(summary)
            top20.sort(key=lambda item: (-item.get("species_count", 0), -item.get("record_count", 0), item.get("location_name", "")))
            return jsonify({"batchNo": result["batchNo"], "top20": top20[:20]})
        return jsonify({"batchNo": "", "top20": []})
    try:
        conn = get_db()
        with conn.cursor() as cur:
            cur.execute("SELECT batch_no FROM bird_record_summary ORDER BY query_time DESC, id DESC LIMIT 1")
            row = cur.fetchone()
            if not row:
                conn.close()
                return jsonify({"batchNo": "", "top20": []})
            batch_no = row[0]
            cur.execute("""
                SELECT id, batch_no, location_name, start_date, end_date, species_count,
                       report_count, record_count, individual_count, query_time
                FROM bird_record_summary
                WHERE batch_no = %s
                ORDER BY species_count DESC, record_count DESC, id ASC
                LIMIT 20
            """, (batch_no,))
            top20 = dict_rows(cur)
            blocked = {name.casefold() for name in get_bird_record_blacklist()}
            if blocked:
                cur.execute("""
                    SELECT summary_id, observation_location, bird_name, report_no, bird_count, record_user
                    FROM bird_record_detail WHERE batch_no = %s
                """, (batch_no,))
                grouped = {}
                for detail in dict_rows(cur):
                    if str(detail.get("record_user") or "").strip().casefold() in blocked:
                        continue
                    grouped.setdefault(detail["summary_id"], []).append(detail)
                filtered_top20 = []
                for row in top20:
                    details = grouped.get(row["id"], [])
                    if not details:
                        continue
                    row["species_count"] = len({item.get("bird_name") for item in details if item.get("bird_name")})
                    row["report_count"] = len({item.get("report_no") for item in details if item.get("report_no")})
                    row["record_count"] = len(details)
                    row["individual_count"] = sum(item.get("bird_count") or 0 for item in details)
                    filtered_top20.append(row)
                top20 = sorted(filtered_top20, key=lambda item: (-item["species_count"], -item["record_count"], item["id"]))[:20]
        conn.close()
        for row in top20:
            for key in ("start_date", "end_date", "query_time"):
                if row.get(key):
                    row[key] = row[key].isoformat() if hasattr(row[key], "isoformat") else str(row[key])
        return jsonify({"batchNo": batch_no, "top20": top20})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/api/bird-records/export/latest', methods=['GET'])
def api_bird_records_export_latest():
    init_bird_records_tables()
    conn = None
    try:
        if _manager.db_enabled:
            conn = get_db()
            export_data = get_latest_bird_batch_export_data(conn)
            conn.close()
            conn = None
        else:
            export_data = get_memory_bird_batch_export_data(_LATEST_MEMORY_BATCH)
        if not export_data:
            message = "查询缓存已过期，请重新查询后在五分钟内下载" if not _manager.db_enabled else "还没有可下载的鸟种记录批次"
            return jsonify({"error": message}), 404
        workbook = build_bird_records_export_workbook(export_data)
        filename = f"bird_records_{export_data['batch_no']}.xlsx"
        return send_file(
            workbook,
            as_attachment=True,
            download_name=filename,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        if conn:
            conn.close()

@app.route('/api/bird-records/clear', methods=['POST'])
def api_bird_records_clear():
    """清空所有已保存的鸟种查询记录。"""
    init_bird_records_tables()
    if not _manager.db_enabled:
        global _LATEST_MEMORY_BATCH
        _LATEST_MEMORY_BATCH = None
        with _CACHE_LOCK:
            for key in list(_CACHE):
                if key.startswith(_BIRD_QUERY_CACHE_PREFIX):
                    _CACHE.pop(key, None)
        return jsonify({"status": "ok", "message": "鸟种查询缓存已清空"})
    conn = None
    try:
        conn = get_db()
        with conn.cursor() as cur:
            # 先清子表，再清汇总表，避免外键约束问题。
            cur.execute("TRUNCATE TABLE bird_record_detail")
            cur.execute("TRUNCATE TABLE bird_record_summary")
            cur.execute("TRUNCATE TABLE bird_record_source_detail")
            cur.execute("TRUNCATE TABLE bird_locations")
        conn.commit()
        conn.close()
        return jsonify({"status": "ok", "message": "鸟种查询记录已清空"})
    except Exception as e:
        try:
            if conn:
                conn.rollback()
                conn.close()
        except Exception:
            pass
        return jsonify({"error": str(e)}), 500

@app.route('/api/bird-records/summary/<int:summary_id>/details', methods=['GET'])
def api_bird_records_summary_details(summary_id):
    init_bird_records_tables()
    keyword = request.args.get("keyword", "").strip()
    if not _manager.db_enabled:
        memory_batch = _LATEST_MEMORY_BATCH
        memory_summary = (memory_batch or {}).get("summaries", {}).get(summary_id)
        if not memory_summary or memory_batch["expires_at"] <= time.monotonic():
            return jsonify({"error": "缓存中的汇总记录已过期，请重新查询"}), 404
        details = [item for item in filter_bird_record_blacklist(memory_summary["details"])
                   if item.get("protection_level") in ("Ⅰ级", "Ⅱ级")]
        if keyword:
            details = [item for item in details if keyword.lower() in item.get("bird_name", "").lower()]
        return jsonify({"summary": memory_summary["summary"], "details": details,
                        "birdNames": sorted({item.get("bird_name") for item in details if item.get("bird_name")})})
    try:
        conn = get_db()
        with conn.cursor() as cur:
            cur.execute("""
                SELECT id, batch_no, location_name, start_date, end_date, species_count,
                       report_count, record_count, individual_count, query_time
                FROM bird_record_summary
                WHERE id = %s
            """, (summary_id,))
            summary_row = cur.fetchone()
            if not summary_row:
                conn.close()
                return jsonify({"error": "汇总记录不存在"}), 404
            summary_columns = [desc[0] for desc in cur.description]
            summary = dict(zip(summary_columns, summary_row))
            params = [summary_id]
            where_keyword = ""
            if keyword:
                where_keyword = " AND d.bird_name LIKE %s"
                params.append(f"%{keyword}%")
            cur.execute(f"""
                SELECT d.report_no, d.observation_location, d.bird_name,
                       COALESCE(d.protection_level, p.protection_level) AS protection_level,
                       d.bird_count, d.record_user, d.observation_time, d.outside_type
                FROM bird_record_detail d
                LEFT JOIN protected_wildlife_catalog p ON p.chinese_name = d.bird_name
                WHERE d.summary_id = %s
                  AND COALESCE(d.protection_level, p.protection_level) IN ('Ⅰ级', 'Ⅱ级')
                  {where_keyword}
                ORDER BY d.bird_name ASC, d.report_no ASC, d.id ASC
                LIMIT 2000
            """, tuple(params))
            details = filter_bird_record_blacklist(dict_rows(cur))
            cur.execute("""
                SELECT DISTINCT bird_name
                FROM bird_record_detail
                WHERE summary_id = %s
                ORDER BY bird_name ASC
            """, (summary_id,))
            all_bird_names = [row[0] for row in cur.fetchall() if row[0]]
        conn.close()
        for key in ("start_date", "end_date", "query_time"):
            if summary.get(key):
                summary[key] = summary[key].isoformat() if hasattr(summary[key], "isoformat") else str(summary[key])
        # fetchall() 已经在上面转换成了字符串列表，不能再次按下标取首字符。
        # 否则“白鹭”会被错误返回为“白”，前端摘要就会出现逐字拆分的鸟名。
        all_bird_names = [str(name).strip() for name in all_bird_names if str(name).strip()]
        return jsonify({"summary": summary, "details": details, "birdNames": all_bird_names})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/api/bird-records', methods=['GET'])
def api_bird_records():
    location = request.args.get("location", "").strip()
    today = datetime.now()
    default_start = (today - timedelta(days=1)).strftime("%Y-%m-%d")
    start_date = request.args.get("startDate", "").strip() or request.args.get("date", "").strip() or default_start
    end_date = request.args.get("endDate", "").strip() or today.strftime("%Y-%m-%d")
    if not location:
        return jsonify({"error": "缺少地点"}), 400
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", start_date) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", end_date):
        return jsonify({"error": "日期格式必须为 YYYY-MM-DD"}), 400

    cache_key = _query_cache_key(_BIRD_QUERY_CACHE_PREFIX, {
        "location": location, "start_date": start_date, "end_date": end_date,
    })
    cached = _cache_get(cache_key)
    if cached is not None:
        return jsonify(cached)

    base_params = {
        "startTime": start_date,
        "endTime": end_date,
        "pointname": location,
        "version": "CH4",
        "mode": "0",
        "page": "1",
        "limit": "1500",
    }
    try:
        record_groups = []
        raw_record_count = 0
        for outside_type in (0, 1):
            params = dict(base_params, outside_type=str(outside_type))
            response = birdreport_post("front/record/activity/taxon", params)
            if response.get("code") not in (0, "0"):
                return jsonify({"error": response.get("msg") or "birdreport 查询失败"}), 502
            raw_record_count += int(response.get("count") or 0)
            records = decode_birdreport_data(response.get("data")) if response.get("count", 0) else []
            for record in records:
                record["outside_type"] = outside_type
            record_groups.append((outside_type, records))
        records, total_records = merge_bird_taxon_records(record_groups)
        result = {
            "location": location,
            "startDate": start_date,
            "endDate": end_date,
            "count": len(records),
            "recordTotal": total_records or raw_record_count,
            "records": records,
            "source": "https://www.birdreport.cn/home/search/page.html"
        }
        return jsonify(_cache_set(cache_key, result))
    except urllib.error.HTTPError as e:
        return jsonify({"error": f"birdreport 返回 {e.code}"}), 502
    except urllib.error.URLError as e:
        return jsonify({"error": f"无法连接 birdreport: {e.reason}"}), 502
    except Exception as e:
        return jsonify({"error": str(e)}), 500

def _deepseek_config():
    cfg = _manager.get_all().get("deepseek") or {}
    try:
        timeout = int(cfg.get("timeout") or 90)
    except (TypeError, ValueError):
        timeout = 90
    return {
        "model": str(cfg.get("model") or "deepseek-chat").strip(),
        "api_key": str(cfg.get("api_key") or "").strip(),
        "base_url": str(cfg.get("base_url") or "https://api.deepseek.com").strip().rstrip("/"),
        "timeout": max(10, min(timeout, 180)),
    }


_AI_CITY_ALIASES = {
    "广州市": ("广东省", "广州市"), "广州": ("广东省", "广州市"),
    "深圳市": ("广东省", "深圳市"), "深圳": ("广东省", "深圳市"),
    "杭州市": ("浙江省", "杭州市"), "杭州": ("浙江省", "杭州市"),
    "青岛市": ("山东省", "青岛市"), "青岛": ("山东省", "青岛市"),
    "南京市": ("江苏省", "南京市"), "南京": ("江苏省", "南京市"),
    "成都市": ("四川省", "成都市"), "成都": ("四川省", "成都市"),
    "武汉市": ("湖北省", "武汉市"), "武汉": ("湖北省", "武汉市"),
    "西安市": ("陕西省", "西安市"), "西安": ("陕西省", "西安市"),
    "厦门市": ("福建省", "厦门市"), "厦门": ("福建省", "厦门市"),
    "北京市": ("北京市", ""), "北京": ("北京市", ""),
    "上海市": ("上海市", ""), "上海": ("上海市", ""),
    "天津市": ("天津市", ""), "天津": ("天津市", ""),
    "重庆市": ("重庆市", ""), "重庆": ("重庆市", ""),
}
_AI_LANDMARK_ALIASES = {
    "越秀公园": ("广东省", "广州市"),
    "海珠湿地": ("广东省", "广州市"),
    "白云山": ("广东省", "广州市"),
}


def _extract_ai_location(question):
    """从原始问题提取明确城市，避免模型把偏好地点带入本次查询。"""
    text = re.sub(r"\s+", "", str(question or ""))
    for alias in sorted(_AI_LANDMARK_ALIASES, key=len, reverse=True):
        if alias in text:
            return _AI_LANDMARK_ALIASES[alias]
    for alias in sorted(_AI_CITY_ALIASES, key=len, reverse=True):
        if alias in text:
            return _AI_CITY_ALIASES[alias]
    province_aliases = {
        "广东": "广东省", "浙江": "浙江省", "山东": "山东省", "江苏": "江苏省",
        "四川": "四川省", "湖北": "湖北省", "福建": "福建省", "陕西": "陕西省",
        "河北": "河北省", "河南": "河南省", "云南": "云南省", "海南": "海南省",
    }
    for alias, province in province_aliases.items():
        if alias in text:
            return province, ""
    return "", ""


def _extract_ai_landmark(question):
    text = re.sub(r"\s+", "", str(question or ""))
    for alias in sorted(_AI_LANDMARK_ALIASES, key=len, reverse=True):
        if alias in text:
            return alias
    return ""


def _requires_ai_bird_records(question):
    value = re.sub(r"\s+", "", str(question or ""))
    time_sensitive = re.search(r"最近|近期|今天|昨日|昨天|前天|这两天|这几天|近[一二三四五六七八九十\d]+天|本周|这周|本月|这个月|今年", value)
    bird_context = re.search(r"鸟|观测|观察|记录|鸟讯|鸟况|值得看|能看到|去哪看|去哪里看", value)
    live_question = re.search(r"有什么鸟|有哪些鸟|哪些鸟|哪里有鸟|哪里能看|去哪看鸟|去哪里看鸟|能看到什么|值得看什么|鸟讯|鸟况|观鸟记录|观察记录|观测记录|记录最多|鸟种最多|近期记录|最新记录", value)
    return bool(live_question or (time_sensitive and bird_context))


def _bird_ai_system_prompt(today):
    return (
        f"你是鸟友工具箱的专业观鸟助手。今天是 {today}。\n"
        "【何时查数据】凡是询问近期/指定日期的鸟种、数量、地点、鸟况、鸟讯、哪里值得看或记录排行，必须先调用 "
        "query_bird_records；不得凭常识补充本次查询中没有出现的鸟名或数量。纯鸟类知识、辨识方法、行为习性问题可以直接回答。\n"
        "【地点】必须根据本次问题确定查询地点；用户没有说明地点时不要猜测，要求用户补充。城市必须尽量补全所属省份。\n"
        "【日期优先级】先看本次问题，再结合历史对话补全省略的日期；如果明确说了‘今天’、‘昨天’、‘这两天’、最近N天、本周等范围，必须严格按该范围换算 YYYY-MM-DD；"
        "只有整个对话都没有说明时间时，才使用程序默认的最近7天。\n"
        "【回答原则】先给简短结论，再按地点或鸟种列出最有用的结果；明确实际查询日期、区域、地点数、记录数和数据来源。"
        "‘有记录’不等于现在一定能看到，推荐时说明这是基于近期公开记录。零结果时说明实际条件并建议扩大日期或区域，不得猜测。"
        "使用简洁中文和短列表，不堆砌标题，不使用宽表格。"
    )


def _chinese_number(value):
    digits = {"一": 1, "二": 2, "两": 2, "三": 3, "四": 4, "五": 5,
              "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}
    text = str(value or "").strip()
    if text.isdigit():
        return int(text)
    if text in digits:
        return digits[text]
    if len(text) == 2 and text[0] == "十" and text[1] in digits:
        return 10 + digits[text[1]]
    if len(text) == 2 and text[1] == "十" and text[0] in digits:
        return digits[text[0]] * 10
    return None


def _extract_ai_date_range(question, today):
    """从用户原话确定日期，优先级高于模型传回的日期参数。"""
    text = re.sub(r"\s+", "", str(question or ""))
    date_match = re.search(r"(20\d{2})[-年](\d{1,2})[-月](\d{1,2})日?(?:至|到|~|～|-)(20\d{2})?[-年]?(\d{1,2})[-月](\d{1,2})日?", text)
    if date_match:
        try:
            start = datetime(int(date_match.group(1)), int(date_match.group(2)), int(date_match.group(3)))
            end = datetime(int(date_match.group(4) or date_match.group(1)), int(date_match.group(5)), int(date_match.group(6)))
            return start, end, "user"
        except ValueError:
            pass
    single_date = re.search(r"(20\d{2})[-年](\d{1,2})[-月](\d{1,2})日?", text)
    if single_date:
        try:
            day = datetime(int(single_date.group(1)), int(single_date.group(2)), int(single_date.group(3)))
            return day, day, "user"
        except ValueError:
            pass
    if re.search(r"今天", text):
        return today, today, "user"
    if re.search(r"昨天|昨日", text):
        day = today - timedelta(days=1)
        return day, day, "user"
    if re.search(r"前天", text):
        day = today - timedelta(days=2)
        return day, day, "user"
    match = re.search(r"(?:最近|近|过去|这)([一二两三四五六七八九十\d]+)天", text)
    if match:
        days = _chinese_number(match.group(1))
        if days and days > 0:
            return today - timedelta(days=days - 1), today, "user"
    if "本周" in text or "这周" in text:
        start = today - timedelta(days=today.weekday())
        return start, today, "user"
    if "本月" in text or "这个月" in text:
        return today.replace(day=1), today, "user"
    if "今年" in text:
        return today.replace(month=1, day=1), today, "user"
    return None


def _render_ai_markdown(value):
    """将 AI 回复转换为安全的 Markdown HTML，供桌面端对话气泡展示。"""
    source = html_escape(str(value or ""), quote=False)
    rendered = markdown.markdown(
        source,
        extensions=["fenced_code", "tables", "nl2br", "sane_lists"],
    )
    # Markdown 链接只允许常见安全协议；异常协议降级为不可点击链接。
    rendered = re.sub(
        r'(<a\s+[^>]*href=")([^"#]+)(")',
        lambda match: match.group(1) + (
            match.group(2)
            if re.match(r"^(?:https?|mailto):", match.group(2), re.IGNORECASE)
            else "#"
        ) + match.group(3),
        rendered,
        flags=re.IGNORECASE,
    )
    rendered = re.sub(r'<a\s+', '<a target="_blank" rel="noopener noreferrer" ', rendered)
    return rendered


def _ai_bird_records_tool(arguments, question="", conversation_context=""):
    today = datetime.now()
    # 当前问题优先；追问省略地点时，再从历史对话补全。
    explicit_province, explicit_city = _extract_ai_location(question)
    explicit_landmark = _extract_ai_landmark(question)
    if not explicit_province and conversation_context:
        explicit_province, explicit_city = _extract_ai_location(conversation_context)
        explicit_landmark = _extract_ai_landmark(conversation_context)
    explicit_question_location = bool(explicit_province)
    if explicit_question_location:
        province, city = explicit_province, explicit_city
        district = str(arguments.get("district") or "").strip() if explicit_city else ""
        location_keyword = explicit_landmark or str(arguments.get("location_keyword") or "").strip()
    else:
        province = str(arguments.get("province") or "").strip()
        city = str(arguments.get("city") or "").strip()
        district = str(arguments.get("district") or "").strip()
        location_keyword = str(arguments.get("location_keyword") or "").strip()
    # 日期同样优先当前问题；“那这两天呢”会从当前问题得到两天范围。
    detected_dates = _extract_ai_date_range(question, today)
    if not detected_dates and conversation_context:
        detected_dates = _extract_ai_date_range(conversation_context, today)
    if detected_dates:
        start_day, end_day, date_source = detected_dates
        start_date = start_day.strftime("%Y-%m-%d")
        end_date = end_day.strftime("%Y-%m-%d")
    else:
        # 没有明确时间时固定使用最近 7 天；不采信模型可能误填的日期。
        end_date = today.strftime("%Y-%m-%d")
        start_date = (today - timedelta(days=6)).strftime("%Y-%m-%d")
        date_source = "recent_7_days"
    validate_bird_dates(start_date, end_date)
    if not province:
        raise ValueError("未指定查询地点，请在问题中说明省份、城市、公园或其他地点")

    details = filter_bird_record_blacklist(fetch_birdreport_public_area_details(province, start_date, end_date, city))
    levels = get_protected_wildlife_level({item.get("bird_name") for item in details})
    for item in details:
        item["protection_level"] = levels.get(item.get("bird_name"))
    for filter_value in (district, location_keyword):
        if filter_value:
            keyword = filter_value.casefold()
            details = [item for item in details if keyword in (item.get("observation_location") or "").casefold()]

    grouped = group_bird_details_by_location(details)
    locations = []
    for item in grouped:
        item_details = item["details"]
        bird_names = sorted(
            {d.get("bird_name") for d in item_details if d.get("bird_name")},
            key=lambda name: (0 if levels.get(name) == "Ⅰ级" else 1 if levels.get(name) == "Ⅱ级" else 2, name),
        )
        bird_details = []
        for name in bird_names:
            matching = next((d for d in item_details if d.get("bird_name") == name), {})
            bird_details.append({"name": name, "protection_level": levels.get(name) or "", "uncertain": bool(matching.get("bird_uncertain"))})
        locations.append({
            "location": item["location"],
            "species_count": item["species_count"],
            "record_count": item["record_count"],
            "report_count": item["report_count"],
            "protected_count": sum(1 for name in bird_names if levels.get(name) in ("Ⅰ级", "Ⅱ级")),
            "bird_names": bird_names,
            "bird_details": bird_details,
        })
    locations.sort(key=lambda item: (-item["protected_count"], -item["species_count"], -item["record_count"], item["location"]))
    locations = locations[:20]
    return {
        "query": {
            "province": province,
            "city": city,
            "district": district,
            "location_keyword": location_keyword,
            "start_date": start_date,
            "end_date": end_date,
            "location_source": "user",
            "date_source": date_source,
            "fixed_recent_days": 7 if date_source == "recent_7_days" else None,
        },
        "record_total": len(details),
        "location_total": len(grouped),
        "locations": locations,
        "source": "观鸟数据中心 https://www.birdreport.cn/home/search/page.html",
    }


def _call_deepseek(messages, force_records_query=False):
    cfg = _deepseek_config()
    if not cfg["api_key"]:
        raise RuntimeError("尚未配置 DeepSeek API Key，请先在配置管理中填写")
    payload = {
        "model": cfg["model"],
        "messages": messages,
        "temperature": 0.2,
        "tools": [{
            "type": "function",
            "function": {
                "name": "query_bird_records",
                "description": "查询观鸟数据中心在指定日期、区域或地点关键词下的公开鸟种记录。",
                "parameters": {
                    "type": "object",
                    "properties": {
                        "location_source": {"type": "string", "enum": ["user"], "description": "始终使用本次问题中的地点"},
                        "date_source": {"type": "string", "enum": ["user", "recent_7_days"], "description": "仅作记录，程序会根据用户原话最终确定日期"},
                        "start_date": {"type": "string", "description": "可填写 YYYY-MM-DD，但程序优先解析用户原话"},
                        "end_date": {"type": "string", "description": "可填写 YYYY-MM-DD，但程序优先解析用户原话"},
                        "province": {"type": "string", "description": "本次查询省或直辖市"},
                        "city": {"type": "string", "description": "本次查询城市；直辖市可留空"},
                        "district": {"type": "string", "description": "用户说了区县时填写，否则留空"},
                        "location_keyword": {"type": "string", "description": "地点关键词，例如 天坛；没有地点限制时留空"},
                    },
                    "required": ["location_source", "date_source", "start_date", "end_date", "province", "city", "district", "location_keyword"],
                    "additionalProperties": False,
                },
            },
        }],
        "tool_choice": ({"type": "function", "function": {"name": "query_bird_records"}}
                        if force_records_query else "auto"),
    }
    response = None
    retryable_statuses = {429, 500, 502, 503, 504}
    for attempt in range(3):
        try:
            response = requests.post(
                f"{cfg['base_url']}/chat/completions",
                headers={"Authorization": f"Bearer {cfg['api_key']}", "Content-Type": "application/json"},
                json=payload,
                timeout=cfg["timeout"],
            )
        except requests.RequestException:
            if attempt >= 2:
                raise
            delay = 2 ** attempt
            app.logger.warning("DeepSeek 网络异常，将在 %s 秒后重试（第 %s/3 次）", delay, attempt + 1)
            time.sleep(delay)
            continue
        if response.status_code not in retryable_statuses or attempt >= 2:
            break
        retry_after = response.headers.get("Retry-After", "")
        try:
            delay = max(1, min(int(float(retry_after)), 10)) if retry_after else 2 ** attempt
        except ValueError:
            delay = 2 ** attempt
        app.logger.warning(
            "DeepSeek 返回 HTTP %s，将在 %s 秒后重试（第 %s/3 次）",
            response.status_code, delay, attempt + 1,
        )
        time.sleep(delay)
    if response is None:
        raise RuntimeError("DeepSeek 请求没有收到响应")
    if response.status_code >= 400:
        try:
            payload = response.json()
        except ValueError:
            payload = None
        error_value = payload.get("error") if isinstance(payload, dict) else None
        if isinstance(error_value, dict):
            detail = error_value.get("message") or error_value.get("detail") or response.text
        elif error_value:
            detail = str(error_value)
        else:
            detail = response.text
        if response.status_code == 503:
            raise RuntimeError(f"DeepSeek 服务当前繁忙，已自动重试 3 次，请稍后再试：{detail[:500]}")
        raise RuntimeError(f"DeepSeek 请求失败（{response.status_code}）：{detail[:500]}")
    return response.json()


@app.route('/api/ai-bird-chat', methods=['POST'])
def api_ai_bird_chat():
    data = request.get_json(silent=True) or {}
    user_message = str(data.get("message") or "").strip()
    if not user_message:
        return jsonify({"error": "请输入问题"}), 400
    history = []
    for item in (data.get("messages") or [])[-10:]:
        if isinstance(item, dict) and item.get("role") in ("user", "assistant") and item.get("content"):
            history.append({"role": item["role"], "content": str(item["content"])[:4000]})
    conversation_context = "\n".join(item["content"] for item in history)
    today = datetime.now().strftime("%Y-%m-%d")
    messages = [{
        "role": "system",
        "content": _bird_ai_system_prompt(today),
    }]
    messages.extend(history)
    messages.append({"role": "user", "content": user_message})
    try:
        tool_result = None
        records_required = (
            _requires_ai_bird_records(user_message)
            or _requires_ai_bird_records(conversation_context)
        )
        tool_attempted = False
        for _ in range(3):
            result = _call_deepseek(messages, force_records_query=records_required and tool_result is None and not tool_attempted)
            choice = (result.get("choices") or [{}])[0]
            assistant_message = choice.get("message") or {}
            tool_calls = assistant_message.get("tool_calls") or []
            if not tool_calls:
                if records_required and tool_result is None:
                    return jsonify({"error": "此问题需要查询真实记录，但未能完成查询；请在问题中明确地点和时间范围"}), 400
                answer = assistant_message.get("content") or "DeepSeek 没有返回文字答案。"
                return jsonify({
                    "answer": answer,
                    "answer_html": _render_ai_markdown(answer),
                    "tool_result": tool_result,
                    "records_queried": bool(tool_result),
                    "query_summary": (
                        f"{tool_result.get('query', {}).get('start_date', '')} 至 "
                        f"{tool_result.get('query', {}).get('end_date', '')} · "
                        f"{tool_result.get('query', {}).get('province', '')}"
                        f"{tool_result.get('query', {}).get('city', '')} · "
                        f"{tool_result.get('record_total', 0)} 条记录"
                    ) if tool_result else "",
                    "location_details": (tool_result or {}).get("locations", []),
                })
            messages.append(assistant_message)
            for tool_call in tool_calls:
                function = tool_call.get("function") or {}
                if function.get("name") != "query_bird_records":
                    continue
                try:
                    tool_attempted = True
                    arguments = json.loads(function.get("arguments") or "{}")
                    tool_result = _ai_bird_records_tool(arguments, user_message, conversation_context)
                    tool_content = json.dumps(tool_result, ensure_ascii=False)
                except (ValueError, TypeError) as exc:
                    tool_content = json.dumps({"error": str(exc)}, ensure_ascii=False)
                messages.append({"role": "tool", "tool_call_id": tool_call.get("id"), "content": tool_content[:30000]})
        return jsonify({"error": "DeepSeek 多次调用工具仍未完成回答"}), 502
    except BirdreportCaptchaRequired as exc:
        return jsonify({"error": f"鸟种数据查询需要验证码，请先在鸟种记录页面完成验证：{exc}", "captchaRequired": True}), 429
    except (requests.RequestException, RuntimeError, ValueError) as exc:
        return jsonify({"error": str(exc)}), 502


@app.route('/api/bird-navigation/suggest', methods=['GET'])
def api_bird_navigation_suggest():
    keyword = request.args.get("keyword", "").strip()
    if not keyword:
        return jsonify([])
    if len(keyword) > 50:
        keyword = keyword[:50]
    try:
        if not _manager.db_enabled:
            keyword_lower = keyword.lower()
            rows = [row for row in load_protected_wildlife_cache()
                    if keyword_lower in row.get("chinese_name", "").lower()
                    and row.get("protection_level") in ("Ⅰ级", "Ⅱ级")][:20]
            return jsonify(rows)
        conn = get_db()
        with conn.cursor() as cur:
            cur.execute("""
                SELECT chinese_name, latin_name, protection_level
                FROM protected_wildlife_catalog
                WHERE protection_level IN ('Ⅰ级', 'Ⅱ级')
                  AND chinese_name LIKE %s
                ORDER BY chinese_name ASC
                LIMIT 20
            """, (f"%{keyword}%",))
            rows = dict_rows(cur)
        conn.close()
        return jsonify(rows)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

def fetch_bird_navigation_records(bird_name, province, city, district, start_date, end_date):
    """查询区域内公开记录，并按鸟名筛选。"""
    records = []
    region_name = build_bird_region_name(province, city)
    for outside_type in (0, 1):
        page = 1
        limit = 1500
        while True:
            params = {
                "startTime": start_date,
                "endTime": end_date,
                "province": province,
                "state": "2",
                "version": "CH4",
                "mode": "0",
                "outside_type": str(outside_type),
                "page": str(page),
                "limit": str(limit),
            }
            if city:
                params["city"] = city
            if district:
                params["district"] = district
            response = birdreport_post("front/record/search/page", params)
            if response.get("code") not in (0, "0"):
                raise_birdreport_error(response, "小鸟导航查询失败")
            total_count = int(response.get("count") or 0)
            if total_count <= 0:
                break
            raw_records = decode_birdreport_data(response.get("data"))
            for raw in raw_records:
                raw_name = str(raw.get("taxon_name") or raw.get("taxonname") or "").strip()
                if bird_name not in raw_name:
                    continue
                raw["outside_type"] = outside_type
                detail = normalize_bird_detail(raw, region_name)
                if not detail:
                    continue
                detail["map_url"] = "https://www.amap.com/search?" + urllib.parse.urlencode({
                    "query": detail["observation_location"]
                })
                records.append(detail)
            if page * limit >= total_count or not raw_records:
                break
            page += 1
    records.sort(key=lambda item: (item.get("observation_time") or "", item.get("report_no") or ""), reverse=True)
    return records

@app.route('/api/bird-navigation/search', methods=['POST'])
def api_bird_navigation_search():
    data = request.get_json(silent=True) or {}
    bird_name = str(data.get("birdName") or "").strip()
    today = datetime.now()
    default_start = (today - timedelta(days=1)).strftime("%Y-%m-%d")
    start_date = str(data.get("startDate") or default_start).strip()
    end_date = str(data.get("endDate") or today.strftime("%Y-%m-%d")).strip()
    province = str(data.get("province") or "北京市").strip()
    city = str(data.get("city") or "").strip()
    district = str(data.get("district") or "").strip()
    if not bird_name:
        return jsonify({"error": "请输入或选择鸟种"}), 400
    if not province:
        return jsonify({"error": "省份不能为空"}), 400
    try:
        validate_bird_dates(start_date, end_date)
        cache_key = _query_cache_key(_BIRD_QUERY_CACHE_PREFIX, {
            "bird_name": bird_name, "province": province, "city": city,
            "district": district, "start_date": start_date, "end_date": end_date,
        })
        cached = _cache_get(cache_key)
        if cached is not None:
            return jsonify(cached)
        protection_level = get_protected_wildlife_level([bird_name]).get(bird_name)
        records = fetch_bird_navigation_records(
            bird_name, province, city, district, start_date, end_date
        )
        for record in records:
            record["protection_level"] = protection_level
        result = {
            "birdName": bird_name,
            "province": province,
            "city": city,
            "district": district,
            "startDate": start_date,
            "endDate": end_date,
            "protectionLevel": protection_level,
            "count": len(records),
            "records": records,
        }
        return jsonify(_cache_set(cache_key, result))
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except BirdreportCaptchaRequired as e:
        return jsonify({"error": str(e), "captchaRequired": True}), 429
    except Exception as e:
        return jsonify({"error": str(e)}), 500

# ---- 资金调拨流程测试 ----

FUND_TRANSFER_INSERT_SQL = """
INSERT INTO paymentdb.T_CEBANK_QUERY (
    CE_ID,
    SERIAL_NO,
    ACCOUNT_CUR,
    AMOUNT,
    DOC_FLAG,
    SUMMARY,
    CHECK_NUM,
    OP_TELLER,
    TRANS_TIME,
    OTHER_ACCT,
    OTHER_ACCT_NAME,
    USER_REM,
    DATE_NET_SERIAL,
    CREATE_TIME,
    STATUS,
    REMARKS,
    ACCT_NAME,
    ACCT_NO,
    RECEIVE_BANK_NAME,
    BANK_ID
)
VALUES (
    (SELECT NVL(MAX(CE_ID), 0) + 1 FROM paymentdb.T_CEBANK_QUERY),
    TO_CHAR(SYSTIMESTAMP, 'YYYYMMDDHH24MISSFF3') || TRUNC(DBMS_RANDOM.VALUE(100, 999)),
    '人民币',
    100002,
    '贷',
    NULL,
    '20R' || TO_CHAR(SYSTIMESTAMP, 'YYYYMMDDHH24MISSFF3') || TRUNC(DBMS_RANDOM.VALUE(1000, 9999)),
    '--',
    SYSTIMESTAMP,
    '12345678900',
    '武汉何其有幸科技有限公司',
    '充值和下发通道不一致.',
    'G9369wyl001zyf' || TO_CHAR(SYSTIMESTAMP, 'YYYYMMDDHH24MISSFF3') || TRUNC(DBMS_RANDOM.VALUE(1000, 9999)),
    SYSTIMESTAMP,
    '0',
    NULL,
    'Mock银行',
    '6213350207668213010',
    '集团内部公司A',
    '600000'
)
"""

FUND_TRANSFER_BY_ID_SQL = """
SELECT CE_ID, SERIAL_NO, AMOUNT, DOC_FLAG, CHECK_NUM, TRANS_TIME, OTHER_ACCT,
       OTHER_ACCT_NAME, USER_REM, DATE_NET_SERIAL, CREATE_TIME, STATUS,
       ACCT_NAME, ACCT_NO, RECEIVE_BANK_NAME, BANK_ID
FROM paymentdb.T_CEBANK_QUERY
WHERE CE_ID = :ce_id
"""

FUND_TRANSFER_ORDER_RESULT_SQL = """
SELECT t.ORDER_NO, t.CE_ID, t.ORDER_STATUS, t.RECHARGE_STATUS, t.TRANSFER_TYPE,
       t.AMOUNT, b.ID AS BANK_ORDER_PK, b.ORDER_ID, b.TRANS_ORDER_ID,
       b.BANK_ORDER_ID, b.ACC_TYPE, b.STATE AS BANK_ORDER_STATE, b.AMT,
       b.TRANS_CODE, b.TRANS_MSG, b.QUERY_CODE, b.QUERY_MSG,
       b.CREATE_TIME, b.UPDATE_TIME
FROM T_TRANSFER_ORDER t
LEFT JOIN T_BANK_ORDER b ON b.ORDER_ID = t.ORDER_NO
WHERE t.CE_ID IN (
    SELECT CE_ID
    FROM T_CEBANK_QUERY
    WHERE T_CEBANK_QUERY.DATE_NET_SERIAL IN ({date_net_serial_binds})
)
"""

FUND_TRANSFER_REVIEW_RESULT_SQL = """
SELECT ID, ORDER_ID, DATE_NET_SERIAL, ACCOUNT_ID, ACC_TYPE, CHANGE_AMT, FEE,
       STATE, AUDIT_STATE, MEMO, CREATE_TIME, UPDATE_TIME
FROM T_ACC_REVIEW_RECORD
WHERE DATE_NET_SERIAL IN (
    SELECT DATE_NET_SERIAL
    FROM T_CEBANK_QUERY
    WHERE T_CEBANK_QUERY.DATE_NET_SERIAL IN ({date_net_serial_binds})
)
"""

FUND_TRANSFER_DETAIL_RESULT_SQL = """
SELECT *
FROM T_ACC_DETAIL_TRANS
WHERE ORDER_ID IN (
    SELECT ORDER_ID
    FROM T_ACC_REVIEW_RECORD
    WHERE DATE_NET_SERIAL IN ({date_net_serial_binds})
)
"""

def get_fund_transfer_config():
    cfg = _manager.get_all().get("fund_transfer_test", {})
    oracle_cfg = cfg.get("oracle", {})
    return {
        "oracle": {
            "user": oracle_cfg.get("user", ""),
            "password": oracle_cfg.get("password", ""),
            "dsn": oracle_cfg.get("dsn", ""),
            "instant_client_dir": oracle_cfg.get("instant_client_dir", ""),
        },
        "charge_url": cfg.get("charge_url", "http://192.168.88.87:8180/BossMgr/ceBankQuery/charge.html"),
        "login_url": cfg.get("login_url", "http://192.168.88.87:8180/BossMgr/login.jsp"),
    }

def get_service_vue_config():
    cfg = _manager.get_all().get("service_vue_frontend", {})
    return {
        "project_path": cfg.get(
            "project_path",
            "/Users/wangpengjing/sbyProject/tax_pay_ext/bossmgr-new-console",
        ),
        "command": cfg.get("command", "pnpm dev"),
        "page_url": cfg.get("page_url", "http://localhost:8080/BossMgr/new/"),
        "login_page_url": cfg.get(
            "login_page_url",
            cfg.get("login_url", "http://localhost:8080/BossMgr/new/login.html"),
        ),
        "login_action_url": cfg.get(
            "login_action_url",
            "http://localhost:8080/BossMgr/j_spring_security_check",
        ),
        "username": cfg.get("username", "admin@fws.com"),
        "password": cfg.get("password", "1234Qwer!"),
    }

def append_service_vue_log(line):
    timestamp = datetime.now().strftime("%H:%M:%S")
    text = f"[{timestamp}] {line.rstrip()}"
    with SERVICE_VUE_LOCK:
        SERVICE_VUE_LOGS.append(text)
        if len(SERVICE_VUE_LOGS) > SERVICE_VUE_MAX_LOG_LINES:
            del SERVICE_VUE_LOGS[:-SERVICE_VUE_MAX_LOG_LINES]

def read_service_vue_output(process):
    try:
        for line in iter(process.stdout.readline, ""):
            if not line:
                break
            append_service_vue_log(line)
    finally:
        code = process.poll()
        if code is not None:
            append_service_vue_log(f"Vue 服务已退出，退出码 {code}")

def service_vue_is_running():
    return SERVICE_VUE_PROCESS is not None and SERVICE_VUE_PROCESS.poll() is None

@app.route('/api/service-vue/config', methods=['GET'])
def api_service_vue_config():
    cfg = get_service_vue_config()
    return jsonify({
        "project_path": cfg["project_path"],
        "command": cfg["command"],
        "page_url": cfg["page_url"],
        "login_page_url": cfg["login_page_url"],
        "login_action_url": cfg["login_action_url"],
        "username": cfg["username"],
    })

@app.route('/api/service-vue/status', methods=['GET'])
def api_service_vue_status():
    process = SERVICE_VUE_PROCESS
    cfg = get_service_vue_config()
    return jsonify({
        "running": service_vue_is_running(),
        "pid": process.pid if process and process.poll() is None else None,
        "return_code": process.poll() if process else None,
        "config": {
            "project_path": cfg["project_path"],
            "command": cfg["command"],
            "page_url": cfg["page_url"],
            "login_page_url": cfg["login_page_url"],
            "login_action_url": cfg["login_action_url"],
            "username": cfg["username"],
        },
    })

@app.route('/api/service-vue/logs', methods=['GET'])
def api_service_vue_logs():
    since = request.args.get("since", 0)
    try:
        since = max(0, int(since))
    except (TypeError, ValueError):
        since = 0
    with SERVICE_VUE_LOCK:
        logs = SERVICE_VUE_LOGS[since:]
        next_index = len(SERVICE_VUE_LOGS)
    return jsonify({
        "logs": logs,
        "next_index": next_index,
        "running": service_vue_is_running(),
    })

@app.route('/api/service-vue/start', methods=['POST'])
def api_service_vue_start():
    global SERVICE_VUE_PROCESS
    cfg = get_service_vue_config()
    project_path = cfg["project_path"]
    if service_vue_is_running():
        return jsonify({
            "status": "ok",
            "message": "Vue 服务已在运行",
            "pid": SERVICE_VUE_PROCESS.pid,
            "running": True,
        })
    if not os.path.isdir(project_path):
        return jsonify({"error": f"项目目录不存在: {project_path}"}), 404

    append_service_vue_log(f"准备启动: {cfg['command']}")
    append_service_vue_log(f"工作目录: {project_path}")
    try:
        popen_kwargs = {
            "cwd": project_path,
            "stdout": subprocess.PIPE,
            "stderr": subprocess.STDOUT,
            "text": True,
            "bufsize": 1,
        }
        if os.name == "posix":
            popen_kwargs["preexec_fn"] = os.setsid
        SERVICE_VUE_PROCESS = subprocess.Popen(
            ["/bin/zsh", "-lc", cfg["command"]],
            **popen_kwargs,
        )
        threading.Thread(
            target=read_service_vue_output,
            args=(SERVICE_VUE_PROCESS,),
            daemon=True,
        ).start()
        append_service_vue_log(f"已启动 Vue 服务，PID {SERVICE_VUE_PROCESS.pid}")
        return jsonify({
            "status": "ok",
            "message": "已启动 Vue 服务",
            "pid": SERVICE_VUE_PROCESS.pid,
            "running": True,
        })
    except FileNotFoundError as e:
        append_service_vue_log(f"启动失败: {e}")
        return jsonify({"error": "未找到 /bin/zsh 或 pnpm，请检查本机环境"}), 500
    except Exception as e:
        append_service_vue_log(f"启动失败: {e}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/service-vue/stop', methods=['POST'])
def api_service_vue_stop():
    global SERVICE_VUE_PROCESS
    process = SERVICE_VUE_PROCESS
    if not process or process.poll() is not None:
        return jsonify({"status": "ok", "message": "Vue 服务未运行", "running": False})
    try:
        append_service_vue_log(f"正在停止 Vue 服务，PID {process.pid}")
        if os.name == "posix":
            os.killpg(os.getpgid(process.pid), signal.SIGTERM)
        else:
            process.terminate()
        return jsonify({"status": "ok", "message": "已发送停止信号", "running": False})
    except Exception as e:
        append_service_vue_log(f"停止失败: {e}")
        return jsonify({"error": str(e)}), 500

@app.route('/api/service-vue/open', methods=['POST'])
def api_service_vue_open():
    auto_login_url = urllib.parse.urljoin(request.host_url, "api/service-vue/auto-login")
    try:
        webbrowser.open(auto_login_url)
        return jsonify({"status": "ok", "message": "已调用默认浏览器打开页面", "url": auto_login_url})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/api/service-vue/auto-login', methods=['GET'])
def api_service_vue_auto_login():
    cfg = get_service_vue_config()
    return render_template('service_vue_auto_login.html', **cfg)

def get_oracle_driver():
    try:
        import oracledb
        return oracledb
    except ImportError:
        try:
            import cx_Oracle
            return cx_Oracle
        except ImportError:
            return None

def init_oracle_client_if_needed(driver, instant_client_dir):
    global _ORACLE_CLIENT_INITIALIZED
    if _ORACLE_CLIENT_INITIALIZED:
        return
    if not instant_client_dir or not hasattr(driver, "init_oracle_client"):
        return
    driver.init_oracle_client(lib_dir=instant_client_dir)
    _ORACLE_CLIENT_INITIALIZED = True

def normalize_oracle_dsn(dsn):
    dsn = (dsn or "").strip()
    jdbc_prefix = "jdbc:oracle:thin:@"
    if dsn.startswith(jdbc_prefix):
        dsn = dsn[len(jdbc_prefix):]
    # JDBC SID 写法常见为 host:port:sid，python-oracledb/cx_Oracle 使用 host:port/sid 更稳。
    parts = dsn.split(":")
    if len(parts) == 3 and "/" not in parts[2]:
        return f"{parts[0]}:{parts[1]}/{parts[2]}"
    return dsn

def get_oracle_connection():
    driver = get_oracle_driver()
    if driver is None:
        raise RuntimeError("当前 Python 环境缺少 cx_Oracle 或 oracledb，无法连接 Oracle")

    cfg = get_fund_transfer_config()["oracle"]
    if not cfg["user"] or not cfg["password"] or not cfg["dsn"]:
        raise ValueError("请先在配置中心填写 Oracle 用户名、密码和 DSN")

    init_oracle_client_if_needed(driver, cfg.get("instant_client_dir"))
    return driver.connect(
        user=cfg["user"],
        password=cfg["password"],
        dsn=normalize_oracle_dsn(cfg["dsn"]),
    )

def oracle_rows_to_dicts(cursor):
    columns = [desc[0].lower() for desc in cursor.description] if cursor.description else []
    rows = []
    for row in cursor.fetchall():
        item = {}
        for key, value in zip(columns, row):
            if hasattr(value, "isoformat"):
                item[key] = value.isoformat()
            else:
                item[key] = str(value) if value is not None else None
        rows.append(item)
    return rows

def parse_date_net_serials(raw_value):
    values = []
    for item in (raw_value or "").replace("\n", ",").split(","):
        value = item.strip().strip("'").strip('"')
        if value:
            values.append(value)
    return list(dict.fromkeys(values))

def build_in_binds(values, prefix):
    bind_names = [f"{prefix}_{index}" for index in range(len(values))]
    bind_sql = ", ".join(f":{name}" for name in bind_names)
    bind_params = {name: value for name, value in zip(bind_names, values)}
    return bind_sql, bind_params

@app.route('/api/fund-transfer-test/insert-ledger', methods=['POST'])
def api_fund_transfer_insert_ledger():
    conn = None
    try:
        driver = get_oracle_driver()
        if driver is None:
            return jsonify({"error": "当前 Python 环境缺少 cx_Oracle 或 oracledb，无法连接 Oracle"}), 500
        conn = get_oracle_connection()
        cur = conn.cursor()
        ce_id_var = cur.var(getattr(driver, "NUMBER", int))
        cur.execute(FUND_TRANSFER_INSERT_SQL + "\nRETURNING CE_ID INTO :ce_id", ce_id=ce_id_var)
        ce_id = ce_id_var.getvalue()
        if isinstance(ce_id, list):
            ce_id = ce_id[0] if ce_id else None
        if ce_id is not None:
            try:
                ce_id = int(ce_id)
            except (TypeError, ValueError):
                ce_id = str(ce_id)
        conn.commit()
        cur.execute(FUND_TRANSFER_BY_ID_SQL, ce_id=ce_id)
        row = cur.fetchone()
        if row:
            columns = [desc[0].lower() for desc in cur.description] if cur.description else []
            latest = {}
            for key, value in zip(columns, row):
                if hasattr(value, "isoformat"):
                    latest[key] = value.isoformat()
                else:
                    latest[key] = str(value) if value is not None else None
        else:
            latest = {}
        cur.close()
        return jsonify({
            "status": "ok",
            "message": "已新增一条银行流水",
            "ce_id": ce_id,
            "date_net_serial": latest.get("date_net_serial"),
            "latest": latest,
        })
    except Exception as e:
        if conn:
            try:
                conn.rollback()
            except Exception:
                pass
        return jsonify({"error": str(e)}), 500
    finally:
        if conn:
            try:
                conn.close()
            except Exception:
                pass

@app.route('/api/fund-transfer-test/results', methods=['POST'])
def api_fund_transfer_results():
    data = request.get_json(silent=True) or {}
    date_net_serials = parse_date_net_serials(data.get("date_net_serial"))
    if not date_net_serials:
        return jsonify({"error": "请先新增流水，或手工输入 DATE_NET_SERIAL，多个值用逗号分隔"}), 400

    conn = None
    try:
        conn = get_oracle_connection()
        cur = conn.cursor()
        result_sets = {}
        bind_sql, bind_params = build_in_binds(date_net_serials, "dns")
        for key, sql in (
            ("transfer_orders", FUND_TRANSFER_ORDER_RESULT_SQL),
            ("review_records", FUND_TRANSFER_REVIEW_RESULT_SQL),
            ("detail_trans", FUND_TRANSFER_DETAIL_RESULT_SQL),
        ):
            cur.execute(sql.format(date_net_serial_binds=bind_sql), bind_params)
            result_sets[key] = oracle_rows_to_dicts(cur)
        cur.close()
        return jsonify({
            "status": "ok",
            "date_net_serial": ",".join(date_net_serials),
            "date_net_serials": date_net_serials,
            "summary": {
                "transfer_orders": len(result_sets["transfer_orders"]),
                "review_records": len(result_sets["review_records"]),
                "detail_trans": len(result_sets["detail_trans"]),
            },
            "results": result_sets,
        })
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        if conn:
            try:
                conn.close()
            except Exception:
                pass

@app.route('/api/fund-transfer-test/login', methods=['POST'])
def api_fund_transfer_login():
    """自动登录 BossMgr 并返回 Cookie"""
    data = request.json or {}
    username = data.get("username", "admin@fws.com").strip()
    password = data.get("password", "1234Qwer!").strip()

    config = get_fund_transfer_config()
    login_url = config["login_url"]

    session = requests.Session()
    session.headers.update({
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
    })

    try:
        # Step 1: GET 登录页，提取隐藏字段（如 CSRF token）和 Cookie
        page_resp = session.get(login_url, timeout=20)
        page_resp.raise_for_status()

        # 解析隐藏表单字段
        import re as _re
        hidden_fields = {}
        for match in _re.finditer(
            r'<input[^>]*type=["\']hidden["\'][^>]*>',
            page_resp.text, _re.IGNORECASE
        ):
            name_match = _re.search(r'name=["\']([^"\']*)["\']', match.group())
            value_match = _re.search(r'value=["\']([^"\']*)["\']', match.group())
            if name_match:
                hidden_fields[name_match.group(1)] = value_match.group(1) if value_match else ""

        # Step 2: POST 登录
        form_data = {"username": username, "password": password}
        form_data.update(hidden_fields)

        login_resp = session.post(
            login_url,
            data=form_data,
            timeout=20,
        )

        # Step 3: 提取 Cookie（处理多个 Set-Cookie 头）
        raw_cookies = {}
        for cookie in session.cookies:
            raw_cookies[cookie.name] = cookie.value

        cookie_str = "; ".join(f"{k}={v}" for k, v in raw_cookies.items())

        return jsonify({
            "status": "ok",
            "cookies": raw_cookies,
            "cookie_string": cookie_str,
            "login_url": login_url,
            "login_status_code": login_resp.status_code,
        })

    except requests.RequestException as e:
        return jsonify({"error": f"登录请求失败: {str(e)}"}), 502
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route('/api/fund-transfer-test/charge', methods=['POST'])
def api_fund_transfer_charge():
    data = request.json or {}
    cookie = data.get("cookie", "").strip()
    if not cookie:
        return jsonify({"error": "请输入 Cookie，例如 JSESSIONID=..."}), 400

    url = get_fund_transfer_config()["charge_url"]
    try:
        resp = requests.post(
            url,
            headers={
                "Cookie": cookie,
                "User-Agent": "Mozilla/5.0",
                "Content-Type": "application/x-www-form-urlencoded",
            },
            timeout=20,
        )
        preview = resp.text[:3000] if resp.text else ""
        return jsonify({
            "status": "ok" if resp.ok else "failed",
            "status_code": resp.status_code,
            "url": url,
            "body_preview": preview,
        }), 200 if resp.ok else 502
    except requests.RequestException as e:
        return jsonify({"error": str(e), "url": url}), 502

# ---- 系统监控 ----

def format_memory(rss_kb):
    if rss_kb >= 1024 * 1024:
        return f"{rss_kb / 1024 / 1024:.2f} GB"
    return f"{rss_kb / 1024:.1f} MB"

def parse_process_name(command):
    if not command:
        return ""
    app_match = re.search(r"/([^/]+\.app)/Contents/", command)
    if app_match:
        return app_match.group(1)[:-4]
    name = os.path.basename(command.rstrip("/")) or command
    if name.endswith(".app"):
        name = name[:-4]
    return name

def get_application_key(command):
    if not command:
        return "unknown"
    app_match = re.search(r"(.+?/[^/]+\.app)/Contents/", command)
    if app_match:
        return app_match.group(1)
    return parse_process_name(command) or command

def get_direct_application_key(command):
    if not command:
        return None
    app_match = re.search(r"(.+?/[^/]+\.app)/Contents/", command)
    if app_match:
        return app_match.group(1)
    return None

def find_application_owner(process, process_map):
    direct_key = process.get("app_key")
    if direct_key:
        return direct_key

    seen = {process["pid"]}
    parent_pid = process.get("ppid")
    while parent_pid and parent_pid not in seen:
        parent = process_map.get(parent_pid)
        if not parent:
            break
        if parent.get("app_key"):
            return parent["app_key"]
        seen.add(parent_pid)
        parent_pid = parent.get("ppid")
    return None

def get_top_memory_applications(limit=10):
    cmd = ["ps", "-axo", "pid=,ppid=,rss=,pmem=,comm="]
    completed = subprocess.run(cmd, check=True, capture_output=True, text=True, timeout=10)
    process_map = {}
    parsed_processes = []

    for line in completed.stdout.splitlines():
        parts = line.strip().split(maxsplit=4)
        if len(parts) < 5:
            continue
        pid_text, ppid_text, rss_text, mem_percent_text, command = parts
        try:
            pid = int(pid_text)
            ppid = int(ppid_text)
            rss_kb = int(rss_text)
            mem_percent = float(mem_percent_text)
        except ValueError:
            continue
        if pid <= 0 or rss_kb <= 0:
            continue

        process = {
            "pid": pid,
            "ppid": ppid,
            "name": parse_process_name(command),
            "command": command,
            "memory_kb": rss_kb,
            "memory": format_memory(rss_kb),
            "memory_percent": mem_percent,
            "app_key": get_direct_application_key(command),
        }
        parsed_processes.append(process)
        process_map[pid] = process

    applications = {}
    for process in parsed_processes:
        command = process["command"]
        app_key = find_application_owner(process, process_map) or get_application_key(command)
        app = applications.setdefault(app_key, {
            "id": app_key,
            "name": parse_process_name(command) or "未知应用",
            "memory_kb": 0,
            "memory_percent": 0,
            "process_count": 0,
            "pids": [],
            "sample_command": command,
            "processes": [],
        })
        if process.get("app_key") == app_key:
            app["name"] = parse_process_name(command) or app["name"]
            app["sample_command"] = command
        app["memory_kb"] += process["memory_kb"]
        app["memory_percent"] += process["memory_percent"]
        app["process_count"] += 1
        app["pids"].append(process["pid"])
        app["processes"].append(process)

    result = []
    for app in applications.values():
        app["memory"] = format_memory(app["memory_kb"])
        app["memory_percent"] = round(app["memory_percent"], 1)
        app["processes"].sort(key=lambda item: item["memory_kb"], reverse=True)
        app["pids"].sort()
        result.append(app)
    result.sort(key=lambda item: item["memory_kb"], reverse=True)
    return result[:limit]

@app.route('/api/system/processes', methods=['GET'])
def api_system_processes():
    try:
        limit = request.args.get("limit", 10, type=int)
        limit = min(max(limit or 10, 1), 50)
        return jsonify({
            "applications": get_top_memory_applications(limit),
            "updated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        })
    except subprocess.CalledProcessError as e:
        return jsonify({"error": e.stderr.strip() or "获取进程列表失败"}), 500
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/api/system/applications/terminate', methods=['POST'])
def api_terminate_application():
    data = request.json or {}
    pids = data.get("pids", [])
    if not isinstance(pids, list) or not pids:
        return jsonify({"error": "缺少要关闭的 PID 列表"}), 400

    current_pid = os.getpid()
    terminated = []
    errors = []
    for raw_pid in pids:
        try:
            pid = int(raw_pid)
        except (TypeError, ValueError):
            errors.append({"pid": raw_pid, "error": "PID 无效"})
            continue
        if pid <= 0:
            errors.append({"pid": pid, "error": "PID 无效"})
            continue
        if pid == current_pid:
            errors.append({"pid": pid, "error": "不能关闭当前工具箱服务进程"})
            continue
        try:
            os.kill(pid, signal.SIGTERM)
            terminated.append(pid)
        except ProcessLookupError:
            errors.append({"pid": pid, "error": "进程已不存在"})
        except PermissionError:
            errors.append({"pid": pid, "error": "权限不足"})
        except Exception as e:
            errors.append({"pid": pid, "error": str(e)})

    status_code = 200 if terminated else 400
    message = f"已向 {len(terminated)} 个进程发送关闭信号"
    if errors:
        message += f"，{len(errors)} 个进程未关闭"
    return jsonify({
        "status": "ok" if terminated else "failed",
        "message": message,
        "terminated": terminated,
        "errors": errors
    }), status_code

@app.route('/api/system/processes/<int:pid>/terminate', methods=['POST'])
def api_terminate_process(pid):
    if pid <= 0:
        return jsonify({"error": "PID 无效"}), 400
    if pid == os.getpid():
        return jsonify({"error": "不能关闭当前工具箱服务进程"}), 400
    try:
        os.kill(pid, signal.SIGTERM)
        return jsonify({"status": "ok", "message": f"已向 PID {pid} 发送关闭信号"})
    except ProcessLookupError:
        return jsonify({"error": "进程已不存在"}), 404
    except PermissionError:
        return jsonify({"error": "权限不足，无法关闭该进程"}), 403
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/tools/<tool_id>')
def tool_page(tool_id):
    if tool_id == "log-query":
        return render_template('tool_log_query.html', tools=TOOLS)
    if tool_id == "system-monitor":
        return render_template('tool_system_monitor.html', tools=TOOLS)
    if tool_id == "bird-records":
        return render_template('tool_bird_records.html', tools=TOOLS)
    if tool_id == "ai-bird-chat":
        return render_template('tool_ai_bird_chat.html', tools=TOOLS)
    if tool_id == "fund-transfer-test":
        return render_template('tool_fund_transfer_test.html', tools=TOOLS)
    if tool_id == "service-vue":
        return render_template('tool_service_vue.html', tools=TOOLS)
    if tool_id == "photo-classify":
        return render_template('tool_photo_classify.html', tools=TOOLS)
    if tool_id == "bird-navigation":
        return render_template('tool_bird_navigation.html', tools=TOOLS)
    if tool_id == "image-generation":
        return render_template('tool_image_generation.html', tools=TOOLS)
    if tool_id == "git-branch-manager":
        return render_template('tool_git_branch_manager.html', tools=TOOLS)
    if tool_id in TOOL_VIEWER_CONFIG:
        cfg = TOOL_VIEWER_CONFIG[tool_id]
        return render_template(cfg["template"], tools=TOOLS, **cfg)
    # 查找是否是预定义工具
    tool = next((t for t in TOOLS if t["id"] == tool_id), None)
    if tool:
        return render_template('tool_placeholder.html', tool=tool, tools=TOOLS)
    abort(404)

@app.route('/query', methods=['POST'])
def query_logs():
    data = request.get_json(silent=True) or {}
    alias = data.get('alias')
    keyword = (data.get('keyword') or "").strip()
    level = (data.get('level') or "ALL").upper()
    target_date = data.get('date') or datetime.now().strftime("%Y-%m-%d")
    line_count = data.get('line_count', 200)
    try:
        line_count = int(line_count)
    except (TypeError, ValueError):
        return jsonify({"error": "查询行数必须是数字"}), 400
    line_count = max(1, min(line_count, 5000))

    allowed_levels = {"ALL", "DEBUG", "INFO", "WARN", "ERROR"}
    if level not in allowed_levels:
        return jsonify({"error": f"未知的日志级别: {level}"}), 400

    log_mapping = _manager.log_mapping
    if alias not in log_mapping:
        return jsonify({"error": f"未知的模块别名: {alias}"}), 400

    path_prefix, tomcat_dir = log_mapping[alias]
    remote_path = f"/www/epaysch/{tomcat_dir}/logs/catalina.{target_date}.out"
    if keyword:
        remote_cmd = f"grep -F -- {shlex.quote(keyword)} {shlex.quote(remote_path)}"
        if level != "ALL":
            remote_cmd += f" | grep -F -- {shlex.quote(level)}"
        remote_cmd += f" | tail -n {line_count}"
    else:
        if level == "ALL":
            remote_cmd = f"tail -n {line_count} {shlex.quote(remote_path)}"
        else:
            remote_cmd = f"grep -F -- {shlex.quote(level)} {shlex.quote(remote_path)} | tail -n {line_count}"

    ssh = paramiko.SSHClient()
    ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())

    try:
        ssh.connect(**_manager.ssh_config, timeout=10)
        stdin, stdout, stderr = ssh.exec_command(remote_cmd)

        logs = stdout.read().decode('utf-8', errors='replace')
        error = stderr.read().decode('utf-8', errors='replace')

        if "No such file" in error:
            return jsonify({"error": f"服务器文件不存在: {remote_path}"}), 404

        if error.strip():
            return jsonify({"error": f"远程查询失败: {error.strip()}"}), 500

        if not logs.strip():
            if keyword:
                level_text = "全部级别" if level == "ALL" else level
                return jsonify({"error": f"在 {alias} 的 {level_text} 日志中未匹配到关键词"}), 404
            return jsonify({"error": f"在 {alias} 的日志中未找到内容"}), 404

        os.makedirs(LOCAL_LOG_DIR, exist_ok=True)
        current_time = datetime.now().strftime("%H%M")
        filename = f"{alias}_{target_date}_{current_time}.log"
        filepath = os.path.join(LOCAL_LOG_DIR, filename)

        with open(filepath, "w", encoding="utf-8") as f:
            f.write(logs)

        return jsonify({
            "status": "success",
            "count": len(logs.splitlines()),
            "preview": logs[:5000],
            "filename": filename,
            "level": level,
            "line_count": line_count,
        })

    except paramiko.AuthenticationException:
        return jsonify({"error": "SSH 认证失败，请检查配置中心里的 SSH 用户名或密码"}), 500
    except paramiko.ssh_exception.NoValidConnectionsError as e:
        return jsonify({"error": f"SSH 连接失败，请检查服务器地址和端口: {e}"}), 500
    except (paramiko.SSHException, socket.timeout, TimeoutError) as e:
        return jsonify({"error": f"SSH 查询异常: {e}"}), 500
    except Exception as e:
        app.logger.exception("查询远程日志失败")
        return jsonify({"error": str(e)}), 500
    finally:
        ssh.close()

@app.route('/get_latest_log', methods=['GET'])
def get_latest_log():
    try:
        log_dir = LOCAL_LOG_DIR
        if not os.path.exists(log_dir):
            return jsonify({"error": "本地还没有生成任何日志文件"}), 404
        list_of_files = glob.glob(os.path.join(log_dir, '*.log'))
        if not list_of_files:
            return jsonify({"error": "logs 目录为空"}), 404
        latest_file = max(list_of_files, key=os.path.getmtime)
        filename = os.path.basename(latest_file)
        with open(latest_file, 'r', encoding='utf-8', errors='replace') as f:
            content = f.read()
        return jsonify({"filename": filename, "content": content})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route('/api/docs/tree')
def docs_tree():
    return make_tree(DOCS_BASE_PATH, "feature-doc")


@app.route('/api/docs/view')
def docs_view():
    return make_md_view(DOCS_BASE_PATH)


def resolve_markdown_path(base_path, rel_path):
    full_path = os.path.normpath(os.path.join(base_path, rel_path))
    if not full_path.startswith(os.path.normpath(base_path)):
        raise PermissionError("非法路径")
    if not os.path.isfile(full_path) or not full_path.endswith(".md"):
        raise FileNotFoundError("文件不存在或不是 Markdown 文件")
    return full_path


def make_md_view(base_path):
    """生成 markdown 视图的通用函数，复用 docs_view 逻辑"""
    rel_path = request.args.get("path", "")
    try:
        full_path = resolve_markdown_path(base_path, rel_path)
    except PermissionError as e:
        return jsonify({"error": str(e)}), 403
    except FileNotFoundError:
        return jsonify({"error": "文件不存在或不是 Markdown 文件"}), 404
    try:
        with open(full_path, "r", encoding="utf-8") as f:
            md_content = f.read()
        md_content = re.sub(
            r'```mermaid\n(.*?)```',
            lambda m: f'<div class="mermaid">\n{m.group(1).strip()}\n</div>',
            md_content,
            flags=re.DOTALL
        )
        html = markdown.markdown(
            md_content,
            extensions=["fenced_code", "codehilite", "tables", "toc"]
        )
        return jsonify({"html": html, "name": rel_path})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route('/api/docs/open', methods=['POST'])
def docs_open():
    data = request.json or {}
    rel_path = data.get("path", "").strip()
    try:
        full_path = resolve_markdown_path(DOCS_BASE_PATH, rel_path)
        if os.name == "posix" and subprocess.run(["uname"], capture_output=True, text=True).stdout.strip() == "Darwin":
            subprocess.Popen(["open", full_path])
        elif os.name == "nt":
            os.startfile(full_path)
        else:
            subprocess.Popen(["xdg-open", full_path])
        return jsonify({"status": "ok", "message": "已调用系统默认工具打开文档"})
    except PermissionError as e:
        return jsonify({"error": str(e)}), 403
    except FileNotFoundError:
        return jsonify({"error": "文件不存在或不是 Markdown 文件"}), 404
    except Exception as e:
        return jsonify({"error": str(e)}), 500


def make_tree(base_path, root_name):
    """生成目录树的通用函数"""
    def build_tree(dir_path):
        items = []
        try:
            entries = sorted(os.listdir(dir_path))
        except OSError:
            return items
        for name in entries:
            full = os.path.join(dir_path, name)
            if name.startswith('.'):
                continue
            rel = os.path.relpath(full, base_path)
            if os.path.isdir(full):
                items.append({
                    "name": name, "type": "folder",
                    "path": rel, "children": build_tree(full)
                })
            elif name.endswith(".md"):
                items.append({"name": name, "type": "file", "path": rel})
        return items
    return jsonify({"name": root_name, "type": "folder", "children": build_tree(base_path)})


@app.route('/api/skills/tree')
def skills_tree():
    return make_tree(SKILLS_BASE_PATH, "skills")


@app.route('/api/skills/view')
def skills_view():
    return make_md_view(SKILLS_BASE_PATH)


@app.route('/api/skills/summary')
def skills_summary():
    conn = None
    try:
        conn = get_db()
        with conn.cursor() as cur:
            cur.execute("""
                SELECT folder_name, chinese_name, detail_desc, file_path
                FROM skill_summary
                ORDER BY folder_name
            """)
            rows = cur.fetchall()

        base_path = os.path.normpath(SKILLS_BASE_PATH)
        result = []
        for folder_name, chinese_name, detail_desc, file_path in rows:
            normalized_path = os.path.normpath(file_path or "")
            if normalized_path and not os.path.isfile(normalized_path):
                skill_md_path = os.path.join(os.path.dirname(normalized_path), "SKILL.md")
                if os.path.isfile(skill_md_path):
                    normalized_path = os.path.normpath(skill_md_path)
            rel_path = ""
            if normalized_path.startswith(base_path):
                rel_path = os.path.relpath(normalized_path, base_path)
            result.append({
                "folder_name": folder_name,
                "chinese_name": chinese_name,
                "detail_desc": detail_desc,
                "file_path": file_path,
                "rel_path": rel_path,
            })
        return jsonify(result)
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        if conn:
            try:
                conn.close()
            except Exception:
                pass


@app.route('/api/skills/open', methods=['POST'])
def skills_open():
    data = request.json or {}
    rel_path = data.get("path", "").strip()
    try:
        full_path = resolve_markdown_path(SKILLS_BASE_PATH, rel_path)
        if os.name == "posix" and subprocess.run(["uname"], capture_output=True, text=True).stdout.strip() == "Darwin":
            subprocess.Popen(["open", full_path])
        elif os.name == "nt":
            os.startfile(full_path)
        else:
            subprocess.Popen(["xdg-open", full_path])
        return jsonify({"status": "ok", "message": "已调用系统默认工具打开文档"})
    except PermissionError as e:
        return jsonify({"error": str(e)}), 403
    except FileNotFoundError:
        return jsonify({"error": "文件不存在或不是 Markdown 文件"}), 404
    except Exception as e:
        return jsonify({"error": str(e)}), 500


# ---- 照片分类整理 API ----

@app.route('/api/photo-classify/scan', methods=['POST'])
def photo_classify_scan():
    data = request.get_json(silent=True) or {}
    source_dir = data.get('source_dir', photo_classify.DEFAULT_SOURCE)
    photos, errors = photo_classify.scan_photos(source_dir)
    groups = photo_classify.group_by_date(photos)
    return jsonify({
        'total': len(photos),
        'groups': sorted(groups.values(), key=lambda g: g['date'], reverse=True),
        'photos': photos,
        'errors': errors,
    })


@app.route('/api/photo-classify/in-place', methods=['POST'])
def photo_classify_in_place():
    data = request.get_json(silent=True) or {}
    source_dir = data.get('source_dir', '').strip()
    if not source_dir:
        return jsonify({'error': '源目录不能为空'}), 400
    result, errors = photo_classify.classify_in_place(source_dir)
    if result is None:
        return jsonify({'error': '源目录处理失败', 'errors': errors}), 400
    return jsonify({'result': result, 'errors': errors})


@app.route('/api/photo-classify/execute', methods=['POST'])
def photo_classify_execute():
    data = request.get_json(silent=True) or {}
    source_dir = data.get('source_dir', photo_classify.DEFAULT_SOURCE)
    dest_dir = data.get('dest_dir', '')
    mode = data.get('mode', 'copy')

    if not dest_dir:
        return jsonify({'error': '目标目录不能为空'}), 400
    if mode not in ('copy', 'move'):
        return jsonify({'error': 'mode 必须是 copy 或 move'}), 400

    source_real = os.path.realpath(os.path.expanduser(source_dir))
    dest_real = os.path.realpath(os.path.expanduser(dest_dir))
    if source_real == dest_real:
        return jsonify({'error': '源目录和目标目录不能相同'}), 400
    source_dir = source_real
    dest_dir = dest_real

    result, errors = photo_classify.classify_photos(source_dir, dest_dir, mode)
    if result is None:
        return jsonify({'error': '扫描照片失败', 'errors': errors}), 400

    return jsonify({
        'result': result,
        'errors': errors,
    })


if __name__ == '__main__':
    warmup_caches()
    init_bookmarks_table()
    init_bird_records_tables()
    app.run(debug=True, port=PORT)
